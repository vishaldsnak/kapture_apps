# Builds DMTKapture_Content_Transformation_Checklist.xlsx in this folder (docs/) - the ONLY copy of the checklist
# Every path / link / file transformation the MC conversion applies to a document, side by side
# (sheets 'MME checklist' and 'MME test plan': what differs for the MME market, and how to test it):
# source content -> legacy DMT (InfoManager, Prod) -> DMTKapture (Kapture document) -> Preview.
# Re-run after a change to the conversion:  python docs/build_transformation_checklist.py  (needs openpyxl)
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "DMTKapture_Content_Transformation_Checklist.xlsx")

LIB = "/library/MAZDA"
WD = LIB + "/WIRING_DIAGRAMS/<loc>/<model>/<mt>/<fl>/<mat>"

# ---------------------------------------------------------------------------------------------
# Sheet 1 - Checklist
# ---------------------------------------------------------------------------------------------
CHECK_HEAD = ["Ref", "Area", "Applies to", "Source content (as received)", "Prod - legacy DMT to InfoManager",
              "Kapture document - DMTKapture", "Preview - DMT_PREVIEW", "Changed vs Prod", "Code / config",
              "How to validate", "Result", "Remarks"]

CHECK = [
    # ---- A. paths inside the document content ----
    ["A1", "Inline images <img src>", "SM / OSM HTML and ENT documents (New/Old models)",
     "<img src=\"image/ac5uuw00000002.gif\"> (any relative path; only the file name is kept)",
     LIB + "/SERVICE_MANUALS/<loc>/image/ac5uuw00000002.gif\n(OSM: OTHER_SERVICE_MANUALS; flat per locale - no model folders)",
     "Same as Prod: " + LIB + "/SERVICE_MANUALS/<loc>/image/ac5uuw00000002.gif\n(/library/... sent as is - same rule as Job 10, content.import.inline.image.context empty)",
     "/content" + LIB + "/SERVICE_MANUALS/<loc>/image/ac5uuw00000002.gif\n(okassets.web.context put in front in the preview page only)",
     "No",
     "mc/utils/ConversionUtils.replaceOkAssetsImagesSrcContent; image files copied by XcopyUtil.copyFilesToServer (\"image\")",
     "1) Kapture document source: src starts /library/MAZDA/...\n2) File exists at SERVER_OKASSETS_PHYSICAL_PATH + library/MAZDA/.../image/<name>\n3) Preview: image shows, no 404 in browser network tab",
     "", "http(s):// and # srcs are left untouched. SM/OSM images share ONE image folder per locale - same file name from two models overwrites (legacy behaviour)."],
    ["A2", "Inline images <img src>", "WD HTML / ENT documents",
     "<img src=\"image/xyz.gif\">",
     WD + "/image/xyz.gif",
     "Same as Prod",
     "/content" + WD + "/image/xyz.gif",
     "No", "ConversionUtils.replaceOkAssetsImagesSrcContent (model / manual type / facelift / material folders added for WD)",
     "As A1", "", "Folder names lower-cased."],
    ["A3", "PDF links inside content <a href=\"x.pdf\">", "ENT documents (SM / OSM) only",
     "<a href=\"abc.pdf\">",
     LIB + "/<CH>/<loc>/pdf/abc.pdf  (WD: .../<model>/<mt>/<fl>/<mat>/pdf/abc.pdf)",
     "Same path as Prod - and the file name is now kept when the source href has a folder (pdfs/abc.pdf -> .../pdf/abc.pdf)", "/content + the same path",
     "Fixed", "ConversionUtils.replaceOKAssetsPdfPathsInContent (called for ENT only, not for HTML documents)",
     "Click the PDF link in the preview; check the href in the Kapture document",
     "", "Legacy defect (link ended at .../pdf/) fixed 2026-10-04 - Open points O5."],
    ["A4", "Stylesheet <link> content.css / contents.css", "SM / OSM HTML documents (New/Old models)",
     "<link href=\"./content.css\" rel=\"stylesheet\"> (other <link> tags are dropped)",
     "<link href=\"" + LIB + "/<CH>/<loc>/<model>/<mt>/<fl>/<mat>/html/content.css\"> put in front of the content",
     "Same as Prod: the content.css link is sent (the sanitizer keeps the conversion's own stylesheet links - ConversionUtils.isConversionStyleLink). Final content = <style> blocks + content.css link + body (A8)",
     "Included: <link href=\"/content" + LIB + "/<CH>/<loc>/<model>/<mt>/<fl>/<mat>/html/content.css\"> (verified 2026-10-04)",
     "No",
     "ConversionUtils.readLinkCSSContent / isConversionStyleLink; css files copied by processHTMLDirectoryForOkAssets (\"html\"); KaptureHtmlSanitizer(keep link)",
     "Kapture document source: content.css link before the body; preview: css request 200 at /content/library/.../html/content.css",
     "", "Other <link> tags of the source (print.css ...) are dropped, as legacy. Open points O1 (closed)."],
    ["A5", "Inline <style> blocks", "SM / OSM HTML documents",
     "<style>...</style> in <head>",
     "Put in front of the content; rules 'table {..}' and 'table, th, td {..}' removed",
     "Same as Prod (style tags are not stripped)", "Same as Kapture (verified 2026-10-04)",
     "No", "ConversionUtils.readStyleTagContent", "View source of the Kapture document", "", ""],
    ["A6", "ENT custom stylesheet", "ENT documents (SM / OSM)",
     "(none - added by the conversion)",
     "<link href=\"" + LIB + "/GMS3_CUSTOM/mazda-css-jp.css\"> appended to <body>",
     "SENT: <link href=\"" + LIB + "/GMS3_CUSTOM/mazda-css-jp.css\"> - every other link removed first, this one appended to <body> last, then the body is read (A8); the sanitizer keeps this href",
     "Included: <link href=\"/content" + LIB + "/GMS3_CUSTOM/mazda-css-jp.css\"> (verified 2026-10-04)",
     "No",
     "ConversionUtils.addStyleLinkForEntDocuments / entStyleHref; GMS3_CUSTOM_DIRECTORY + GMS3_MC_CUSTOM_CONTENT_CSS; KaptureHtmlSanitizer(keep link)",
     "Kapture document source ends with the mazda-css-jp.css link; preview: css request 200 at /content/library/MAZDA/GMS3_CUSTOM/mazda-css-jp.css",
     "", "Open points O1 (closed)."],
    ["A7", "Scripts, meta, iframe, object, embed, on* attributes, javascript: URLs", "All content sent to Kapture",
     "As in the source HTML",
     "Sent (meta tags removed in some flows)",
     "Removed by KaptureHtmlSanitizer (kapture.sanitize.strip.tags / .attributes / .all.event.attributes / .url.schemes) - same rules as Job 10",
     "Removed - the preview page is built from the sanitized content (exact replica of what Kapture gets)",
     "YES - OPEN FOR DISCUSSION",
     "kapture/KaptureHtmlSanitizer (applied in KaptureContentService to every text attribute when the payload is built; the same sanitizer builds the preview page)",
     "Kapture document source has no <script>, onclick, javascript: ...; only the conversion's own stylesheet links remain",
     "", "OPEN FOR DISCUSSION (user, 2026-10-04): the sanitizer is GENERIC (Job 10 rules); legacy cleaned per process function (sheet 'Legacy cleaning per process'). Kept as is for now; if Kapture rejects content because of a legacy step, that is taken up with the Kapture team. Open points O7."],
    ["A8", "Body only", "HTML / ENT documents (not OEM)",
     "Full HTML page",
     "Only the inner HTML of <body> is sent. HTML: body read, then content.css link + <style> put in front. ENT: MAZDA css appended to <body>, then body read.",
     "Same as Prod, then sanitized (A7)",
     "The same body content, wrapped in a preview page (title, meta charset, _assets/preview_page.js)",
     "No", "WiringDiagramUtils.readBodyContent; preview/PreviewPageWriter.page", "", "", ""],
    ["A9", "Document title", "All documents",
     "Title from esicat.txt / file", "Special characters replaced, <sup> removed",
     "TITLE: same rule (KaptureArticleBuilder.cleanTitle)", "Tree and header show the same cleaned title",
     "No", "mc/utils/KaptureArticleBuilder.cleanTitle", "Compare title in Kapture, preview tree, transaction report", "",
     "SM/OSM PDF: ATTACHMENT_TITLE carries the RAW title."],

    # ---- B. links ----
    ["B1", "Inner links (document to document)", "SM / OSM HTML and ENT documents",
     "<a href=\"sie_1290337.html\"> / <a href=\"../xx/sie_1290337.html\"> / href=\"javascript:openAnotherTab('sie_1290337.html')\" (ENT: .ent added)",
     "Step 1: href -> source path from the locale folder (ja-JP\\<model>\\...\\sie_1290337.html), looked up by source path.\nMapped: index?page=content&id=SM1234567\nNot mapped: index?page=content&id=",
     "Mapped: content/SM1234567 (LINK_URL=content/, Job 10 DOCUMENT rule)\nNot mapped: content/ (user decision - stays bare)",
     "Mapped: SM1234567.html (+#anchor). Click -> opens in the right pane + tree highlight when the document is in the job; otherwise modal 'Document SM1234567 is not part of this job.'\nNot mapped: dead link (no href)",
     "YES - URL pattern",
     "ConversionUtils.prepareInnerLinkPaths; StartMCConversionImpl inner links pass (re-update of documents whose links map only after their targets exist); PreviewPageWriter.page",
     "1) INNER_LINKS_REPORT.xlsx: mapping status Y/N\n2) Kapture document: href content/<id>\n3) Preview: click in-job link, out-of-job link, unmapped link",
     "", "Lookup strings keep backslashes (DB data is backslash paths)."],
    ["B2", "In-page anchor via openAnotherTab('./frame.html#sie_1')", "SM / OSM (Isuzu pattern 2)",
     "href=\"javascript:openAnotherTab('./frame.html#sie_1290337')\"",
     "#sie_1290337 (not tracked)", "Same as Prod", "Same - jumps inside the page",
     "No", "ConversionUtils.prepareInnerLinkPaths", "", "", ""],
    ["B3", "Other manual links (search links)", "SM / OSM HTML and ENT documents",
     "index?page=result_mc&startover=y&qstr=qstr&fac=CMS-CATEGORY-MAZDA-SERVICE_MANUAL_TYPE.WORKSHOP_MANUAL&question_box=<Tag ID>\n(also page=result_mme and page=result; & or &amp;; relative or full URL)",
     "result_mc / result_mme: passed AS IS, target=_blank added.\npage=result: NOT caught - handled as an inner link and broken.",
     "result_mc?question_box=<Tag ID>  /  result_mme?question_box=<Tag ID>  /  result?question_box=<Tag ID>\n(startover, qstr, fac dropped; #fragment kept; target=_blank kept) - same rules as Job 10",
     "Dead link (no href). Click -> modal 'This is an other manual link: result_mc?question_box=<Tag ID>'",
     "YES",
     "kapture/OtherManualLinkRules + OTHER_MANUAL_LINK_RULES / OTHER_MANUAL_LINK_RULE.<n>.match|target (application.properties); ConversionUtils.prepareInnerLinkPaths",
     "OTHER_MANUAL_LINKS_REPORT.xlsx: INNER LINK SOURCE PATH (as received) + KAPTURE LINK (as sent); Kapture document href; preview modal text",
     "", "Changed 2026-10-04. DB tracking keeps the source link."],
    ["B4", "External links http(s)://", "All HTML content",
     "<a href=\"https://...\">", "Untouched", "Untouched", "Untouched - opens normally", "No", "", "", "", ""],
    ["B5", "Anchors #name", "All HTML content", "<a href=\"#x\">", "Untouched", "Untouched", "Untouched", "No", "", "", "", ""],

    # ---- C. WD JS files in OKAssets ----
    ["C1", "window.js", "WD HTML_MC (html folder, Flash pages)",
     "Quoted html / pdf addresses \"/library/MAZDA/...\"",
     "Copied as is (legacy MC window.js operation writes nothing)",
     "OKAssets copy: every quoted \"/library/ -> \"/content/library/ (byte-preserving, idempotent). Source content is not touched.",
     "Page opens from OKAssets with this window.js - no extra handling",
     "YES",
     "mc/utils/XcopyUtil (html movement) -> WiringDiagramUtils.prefixOkAssetsContextInWindowJS",
     "1) grep the OKAssets copy: no \"/library/ left, \"/content/library/ present\n2) job log: count of addresses changed\n3) WINDOW_JS_REPORT.xlsx",
     "", "Decision 2026-10-03."],
    ["C2", "voltageMAP.js", "WD HTML5 (html5/js), MC",
     "':'/newm_<model>/.../xyz.ent'",
     "dokinfoctr/mazdagms3/index?page=detail_mc&id=SM1234567 (not mapped: ...&id=)",
     "mazdagms3/content/SM1234567 (not mapped: mazdagms3/content/)\n= GMS3_INFOCENTER_APPLICATION_CONTENT + LINK_URL_MC_WINDOW_JS + id",
     "OKAssets file NEVER changed by the preview (mazdagms3 uses the same files). The preview page copy overrides window.open: mazdagms3/content/<id> -> opens the document in the right pane, or modal 'not part of this job'",
     "YES - URL pattern",
     "XcopyUtil (html5 movement) -> WiringDiagramUtils.performVoltageMapJSOperation; PreviewPageWriter.PAGE_SCRIPT_SOURCE",
     "VOLTAGE_MAP_JS_REPORT.xlsx SUCCESS/FAILURE per link; open a voltage link in the preview",
     "", "Only .ent targets are converted."],
    ["C3", "voltagelinkMAP.js", "WD HTML5 (html5/js), MC",
     "As C2", "As C2", "As C2", "As C2", "YES - URL pattern",
     "WiringDiagramUtils.performVoltageMapLinkJSOperation", "VOLTAGE_LINK_MAP_JS_REPORT.xlsx; the viewer's voltage dialog in the preview", "", ""],
    ["C4", "Other html5 js files (not common.js)", "WD HTML5 (html5/js)",
     "\"../conn/\" or \"conn/\"",
     "-> /library/MAZDA/WIRING_DIAGRAMS/<loc>/<model>/<mt>/<fl>/<mat>/html5/conn/",
     "Same as Prod (no /content)", "Read from OKAssets as written",
     "No",
     "WiringDiagramUtils.updateJSContentForHTML5",
     "When a model with a conn folder is processed: browser network tab - no 404 on /library/.../conn/",
     "", "The prod CX-5 sample js has NO conn/ references, so nothing changes today - Open points O4."],

    # ---- D. files carried by the document ----
    ["D1", "PDF attachment", "SM / OSM PDF documents",
     "<mat>_PDF/pdf/abc.pdf",
     "InfoManager attachment; PDF also moved to live/staging folders",
     "ATTACHMENTS { ATTACHMENT_TITLE = raw title, ATTACHMENT = file upload }. No live/staging copy.",
     "Copied to DMT_PREVIEW/<job>/files/<document id>/abc.pdf; right pane shows the PDF; header link 'abc.pdf' opens / downloads it",
     "Carrier only",
     "KaptureArticleBuilder.addPdfAttachments; PreviewBuilder.copyPdf",
     "Kapture: attachment present (a 2xx on the article says nothing about attachments); preview: PDF opens",
     "", "Kapture 500s when the file name contains '..' (shared matrix row 18)."],
    ["D2", "PDF attachment", "WD PDF documents",
     "<mat>/pdf/abc.pdf",
     "InfoManager attachment",
     "ATTACHMENT = file upload (no ATTACHMENTS group, no title). Partial update: only TITLE + ATTACHMENT; the HTML file's attributes stay.",
     "As D1", "Carrier only", "KaptureArticleBuilder.wdAttachment / overlayPartialWiringDiagramUpdate", "As D1", "", ""],
    ["D3", "OKAssets copies", "All channels",
     "Source folders html / html5 / image / pdf / djvu / conn / css / js / img / print ...",
     "Copied under " + LIB + "/<CH>/<loc>/... (folder names lower-cased)",
     "Same, physical root SERVER_OKASSETS_PHYSICAL_PATH (MC Dev /s3kapture/ROOT/resources/sites/KAPTURE/content/)",
     "Served at /content/library/...",
     "Root only", "StartMCConversionImpl.startProcessingOKAssets; XcopyUtil.copyFilesToServer",
     "Files present at the physical path; transaction report OKAssets counts", "", "OEM content: see E4."],
]

# ---------------------------------------------------------------------------------------------
# Sheet 2 - display by document type
# ---------------------------------------------------------------------------------------------
TYPE_HEAD = ["Ref", "Document type", "Identified by", "Kapture document attributes", "Copied to OKAssets",
             "Preview - right pane", "Preview - header links", "Remarks", "Result"]
TYPES = [
    ["E1", "HTML (SM / OSM)", "html file in <mat>_HTML/html, New/Old model",
     "TITLE, CONTENT (body HTML), VTOC_FILENAME + SIE_ID (SM), DJVU_FILE_LOCATION / OASIS_FILE_LOCATION empty",
     "html/image(s) -> " + LIB + "/<CH>/<loc>/image/; html/*.css -> .../<mat>/html/",
     "DMT_PREVIEW/<job>/<document id>.html - offline copy of the content (A1-A9, B1-B5 applied)",
     "Open in new tab", ""],
    ["E2", "ENT (SM / OSM)", ".ent file",
     "As E1 (content parsed from the ENT by EntParsing)", "image / pdf folders",
     "As E1", "Open in new tab", "Custom stylesheet A6."],
    ["E3", "PDF (SM / OSM)", "pdf file in <mat>_PDF",
     "TITLE, CONTENT empty, ATTACHMENTS (D1)", "-",
     "The PDF (files/<document id>/<name>.pdf)", "<name>.pdf", "PDF inline in real Chrome not yet tested."],
    ["E4", "OEM content (SM / OSM)",
     "Model type neither New nor Old (Isuzu, Suzuki, Nissan ...), SM or OSM channel. Never WD.",
     "TITLE + OEM_FILE_LOCATION only - the document is NOT loaded into Kapture.\ne.g. " + LIB + "/OTHER_SERVICE_MANUALS/ja-JP/Isuzu_titan_lh/WM/FL0002/E-W1J-3TC4_HTML/html/si_954130.html",
     "EVERY folder of the material folder, as is, SOURCE NAMES AND CASE KEPT (ja-JP, not ja_jp). Text files in the material folder (esicat, display order, VIN) are not copied.",
     "The OEM page itself, loaded from OKAssets: /content + OEM_FILE_LOCATION. Its own relative links / images work inside OKAssets.",
     "Open in new tab",
     "Links inside an OEM page are the OEM's own relative links - not converted, not routed to the preview tree."],
    ["E5", "WD HTML_MC (Flash pages)", "html file in WD <mat>/html",
     "TITLE, DJVU_FILE_LOCATION = " + WD + "/html/<file>.html, HTML5_FONT_SELECTION. Partial update: only TITLE, DJVU_FILE_LOCATION, HTML5_FONT_SELECTION; the PDF file's ATTACHMENT stays.",
     "whole html folder (swf, window.js C1, conn ...)",
     "The page from OKAssets: /content" + WD + "/html/<file>.html",
     "Open in new tab",
     "The pages are Flash (.swf) - a current browser does not play Flash."],
    ["E6", "WD HTML5", "html file in WD <mat>/html5",
     "TITLE, HTML5_FILE_NAME, HTML5_UPLOAD_DIRECTORY = " + WD + "/, DJVU_FILE_LOCATION = " + WD + "/html5/<file>.html, NAVIGATION, HTML5_FONT_SELECTION",
     "whole html5 folder (css, img, js incl. voltage maps C2/C3, plugin, print, svg)",
     "A copy of the page only: DMT_PREVIEW/<job>/wd/<document id>/<same file name>, with <base href=\"/content" + WD + "/html5/\"> + preview_page.js; everything else loads from OKAssets. Folder per document id - the same page name in two models is safe.",
     "Open in new tab",
     "If the source page is missing, the pane opens the OKAssets page (voltage links then do not route)."],
    ["E7", "WD PDF", "pdf file in WD <mat>/pdf (material folder not _HTML/_XML/_PDF/_DJVU)",
     "TITLE + ATTACHMENT (D2)", "pdf folder", "The PDF", "<name>.pdf", ""],
    ["E8", "DJVU (SM / OSM)", "djvu file in <mat>/djvu",
     "TITLE, DJVU_FILE_LOCATION = " + LIB + "/<CH>/<loc>/<model>/<mt>/<fl>/<mat>/djvu/<name>.html",
     "djvu folder", "The .html page from OKAssets: /content + DJVU_FILE_LOCATION", "Open in new tab",
     "Needs a DjVu viewer in the browser."],
]

# ---------------------------------------------------------------------------------------------
# Sheet 3 - open points
# ---------------------------------------------------------------------------------------------
# Legacy cleaning per process function (open point O7)
PROC_HEAD = ["Ref", "Process function (StartMCConversionImpl)", "Document", "Legacy steps on the content, in order",
             "Legacy leaves AS IS", "Kapture sanitizer removes in addition (current)", "Kapture attributes",
             "Preview", "Remarks"]
PROC = [
    ["P1", "processAHTMLFile", "SM / OSM HTML, New / Old model",
     "1) readLinkCSSContent: only content.css / contents.css links kept, path -> .../<mat>/html/<name>; other <link> dropped\n"
     "2) readStyleTagContent: <style> blocks; 'table {..}' and 'table, th, td {..}' rules removed\n"
     "3) readMetaTagDescriptionValue: meta description read (Toyota flag only)\n"
     "4) replaceOkAssetsImagesSrcContent (A1)\n"
     "5) prepareInnerLinkPaths: inner links (B1, B2), other manual links (B3)\n"
     "6) readBodyContent: inner HTML of <body> only (<head> with title / meta / head scripts dropped)\n"
     "7) content = <style> + content.css link + body",
     "Inside the body: <script>, <iframe>, <object>, <embed>, <meta>, <link>, on* attributes, javascript: URLs (javascript:openAnotherTab is converted in step 5)",
     "All the 'legacy leaves as is' items - the content.css link is kept",
     "CONTENT", "Same content as Kapture (after the sanitizer), /content put in front of /library/ paths", ""],
    ["P2", "processAHTMLFile (OEM)", "SM / OSM HTML, model type not New / Old",
     "Content NOT read. OEM_FILE_LOCATION = web path of the file; the folder is copied as is",
     "The OEM page itself (not sent)", "Nothing (TITLE and OEM_FILE_LOCATION are plain text)",
     "TITLE, OEM_FILE_LOCATION", "The OEM page from OKAssets, as is", ""],
    ["P3", "processAENTFile", "SM / OSM ENT",
     "1) EntParsing.parseEntFile -> HTML\n"
     "2) replaceOkAssetsImagesSrcContent (A1)\n"
     "3) replaceOKAssetsPdfPathsInContent (A3)\n"
     "4) prepareInnerLinkPaths (B1-B3)\n"
     "5) addStyleLinkForEntDocuments: every <link> removed, then the MAZDA css link appended to <body> (removal per user rule 2026-10-04 - same pattern as WiringDiagramUtils.removeStyleTag)\n"
     "6) readBodyContent",
     "Inside the body: <script>, <iframe>, <object>, <embed>, <meta>, on* attributes, javascript: URLs",
     "All the 'legacy leaves as is' items - the MAZDA css link is kept",
     "CONTENT", "As P1", ""],
    ["P4", "processAPDFFile", "SM / OSM / WD PDF",
     "No HTML content; the file becomes the attachment (D1, D2)", "-", "Nothing (TITLE only)",
     "TITLE, ATTACHMENTS / ATTACHMENT", "The PDF", ""],
    ["P5", "processADJVUFile", "DJVU",
     "No content; DJVU_FILE_LOCATION = .../djvu/<name>.html", "-", "Nothing", "TITLE, DJVU_FILE_LOCATION",
     "The page from OKAssets", ""],
    ["P6", "processAHTMLFileForWD", "WD HTML_MC (Flash)",
     "Content not sent. HTML5_FONT_SELECTION = inner HTML of the first <font> of the page; DJVU_FILE_LOCATION = page path",
     "The font HTML as it is", "Tags / on* attributes inside the font HTML (the sanitizer runs on every text attribute)",
     "TITLE, DJVU_FILE_LOCATION, HTML5_FONT_SELECTION", "The page from OKAssets", ""],
    ["P7", "processAHTML5FileForWD", "WD HTML5",
     "Content not sent. HTML5_FONT_SELECTION = first <font> inner HTML; HTML5_FILE_NAME, HTML5_UPLOAD_DIRECTORY, DJVU_FILE_LOCATION (html5 page path), NAVIGATION",
     "The font / navigation HTML as it is", "Tags / on* attributes inside the font and navigation HTML",
     "TITLE, NAVIGATION, HTML5_FILE_NAME, HTML5_UPLOAD_DIRECTORY, HTML5_FONT_SELECTION, DJVU_FILE_LOCATION",
     "Copy of the page (wd/<id>/<name>) - the OKAssets page as is, not sanitized", ""],
    ["P8", "All", "Every document",
     "TITLE: special characters replaced, <sup> removed (KaptureArticleBuilder.cleanTitle)", "-",
     "Tags in the title text", "TITLE", "Tree / header", ""],
]

OPEN_HEAD = ["Ref", "Point", "Effect", "Options", "Owner / decision", "Status"]
MME_OPEN = [
    ["O8", "MME display order staging (gms3_dmt_el_gr_nm_dispord, legacy pattern kept - user decision): every document read its rows with its own SELECT (indexed, but one ~200 ms Router round trip per document).",
     "~200 ms per document on MC Dev (BT-50 WM 5,214 documents = ~17 min).", "Done 2026-10-05: the rows the job stages are also indexed in memory by path; a document takes its rows from there (copies). The staging table is still written and cleaned as before; the SELECT stays as the fallback.", "User", "CLOSED"],
    ["O9", "MME: SCM VIN (M17) and VIN-ML mapping (M18) wrote one row at a time (SELECT + UPDATE / INSERT, SCM also a document id SELECT and a new connection per line).",
     "~600-800 ms per line on MC Dev (MySQL Router).", "Done 2026-10-05: chunked IN-list reads, one CASE / IN UPDATE and one INSERT batch per chunk (db.write.batch.size); same statuses and report.", "User", "CLOSED"],
    ["O10", "MC: MCDocumentManagementDAO.updateInnerLinkMappingStatus used a batched UPDATE (executeBatch, multi-statement under rewriteBatchedStatements).",
     "Router can reject it on MC Dev; one statement per document otherwise.", "Done 2026-10-05: same single CASE + IN UPDATE per chunk as MME (M21).", "User", "CLOSED"],
    ["O12", "MC + MME: carline names of vin.txt / scmvin.txt lines (MMEDocumentManagementDAO.getModelDetails) read with one SELECT per VIN line.",
     "Run 5438 (CX-3 + CX-5): 494 statements; ~200 ms each on MC Dev = ~100 s per such job.", "Done 2026-10-05: one SELECT per chunk of carline codes, matched in memory (same trimmed / lower-case comparison, first row wins). Checked against run 5438 on local kapture_dc: 494 of 494 lines identical (321 named, 173 blank).", "User", "CLOSED"],
    ["O13", "MC + MME (shared): after the category sync, MasterDataSyncTransactionDAO.processMDMItemSyncStatus set MDM_SYNC_STATUS = Y with one UPDATE per MDM row.",
     "~250 ms per row through the MySQL Router on MC Dev: 264 ESI rows = 66 s in run 5617.", "Done 2026-10-05: one UPDATE ... WHERE <id> IN (...) per db.write.batch.size ids, all master data types. Verify on the next MC Dev run: one 'processMDMItemSyncStatus :: ... ids n :: rows updated n' line per type.", "User", "CLOSED"],
    ["O14", "Kapture side (MC Dev 5617, SC Dev 5443): article updates rolled back under load; category create time-out (120 s); ESI name 409s; category API ~2.8 s per call (no bulk category endpoint).",
     "10 + 3 documents not updated (a re-run updates them); master data ~18 min of a 44 min job.", "Sent to Kapture 2026-10-05: VDI_logs/DMT_Logs/Kapture_Issue_Update_Rollback_Category_20261005.html. Names of ESI categories = business data.", "Kapture / business", "OPEN (Kapture)"],
    ["O11", "Automation (AutomationDAO, TransactionDAO) still names GMS3_VC_MME_VIN_DTL_ in upper case.",
     "Fails on Linux MySQL when the automation runs for MME.", "Lower-case them when the automation is migrated (InfoManager feature, out of scope).", "User", "OPEN"],
]
OPEN = [
    ["O1", "Stylesheet links: content.css / contents.css (A4) and the ENT MAZDA css (A6).",
     "Rule (user, 2026-10-04): same as legacy - HTML documents carry their content.css / contents.css link, ENT documents the MAZDA css link (every other link removed first, the MAZDA css appended last). Sent to Kapture and shown in the preview.",
     "Done: ConversionUtils.addStyleLinkForEntDocuments removes all links first; KaptureHtmlSanitizer keeps only ConversionUtils.isConversionStyleLink() hrefs",
     "User", "CLOSED"],
    ["O2", "The preview must be an exact replica of the content sent to Kapture, incl. the ENT MAZDA css.",
     "Preview page = sanitized content (same sanitizer, same kept link), /content put in front of /library/ paths, links made preview links.",
     "Done: PreviewBuilder sanitizes before PreviewPageWriter.page", "User", "CLOSED"],
    ["O3", "Paths inside the Kapture document stay /library/... (A1-A6, E4-E8); only the preview adds /content.",
     "Expected by Kapture (user, 2026-10-04) - same as Job 10 (content.import.inline.image.context empty).",
     "-", "User", "CLOSED"],
    ["O4", "Other html5 js files (not common.js) get absolute /library/.../html5/conn/ paths without /content (C4).",
     "Affects how the WD HTML5 viewer loads its conn files - in Kapture AND in the preview (both read the same OKAssets js). The prod CX-5 sample has no conn folder and no conn/ references, so nothing changes today.",
     "If a model with conn/ references appears: same /content prefix as window.js", "User", "WATCH"],
    ["O5", "Legacy defect in ENT PDF links (A3): a source href with a '/' lost its file name -> .../pdf/.",
     "Broken PDF link in Kapture and preview for such hrefs.",
     "Fixed 2026-10-04 in ConversionUtils.replaceOKAssetsPdfPathsInContent", "User", "CLOSED"],
    ["O7", "OPEN FOR DISCUSSION: the Kapture sanitizer is generic (Job 10 rules); legacy cleaned per process function.",
     "Beyond legacy, the sanitizer removes from the BODY: script / meta / iframe / object / embed / other link tags, on* event attributes and javascript: / vbscript: URLs - for every document type and every text attribute (TITLE, NAVIGATION, HTML5_FONT_SELECTION ...). Legacy left these in the body as they were (only the <head> was dropped). See sheet 'Legacy cleaning per process'.",
     "a) keep the generic sanitizer (current)\nb) follow legacy exactly per process function; any Kapture rejection is taken up with the Kapture team\nUser to discuss and decide - no change until then",
     "User", "OPEN FOR DISCUSSION"],
    ["O6", "SM / OSM images share one image folder per locale (A1).",
     "Same image file name from two models: the later copy overwrites the earlier one.",
     "Accepted - no change (user, 2026-10-04)", "User", "CLOSED"],
]

# ---------------------------------------------------------------------------------------------
# MME - what differs from MC (sheet 'MME checklist') and the test plan (sheet 'MME test plan')
# Every MC row (A1-D3, E1-E8, P1-P8) applies to MME as well, unless an M row says otherwise.
# ---------------------------------------------------------------------------------------------
MME_HEAD = ["Ref", "Area", "Legacy MME (DMT_Optimized, InfoManager)", "DMTKapture MME (now)", "Same as MC?",
            "Code / config", "How to validate", "Result", "Remarks"]

MME_ROWS = [
    # ---- M1. market model ----
    ["M1", "Locales", "14 locales with their own tables (cs_cz de_de el_gr en_uk es_es fi_fi fr_fr it_it nl_nl pl_pl pt_pt ru_ru sv_se tr_tr); each locale an independent InfoManager document",
     "Each locale an INDEPENDENT Kapture document (master identifier) - no master / translation link (user decision 2026-10-05). Kapture locale from the property <xx-yy>.",
     "Yes (MC has one locale)", "StartMMEConversionImpl (kaptureLocale = getProperty(locale.toLowerCase()))",
     "Kapture: the document has no translation parent; locale = the schedule locale", None, None],
    ["M2", "DMT tables per locale", "gms3_dmt_<LOC>_[NM_]{IMDOC, INRLKS, ESICAT, VIN, CAT, VINMAST, CVC, DISPORD, CD_DATA} built in UPPER case",
     "Same tables, names lower case (Linux MySQL is case-sensitive): gms3_dmt_<loc>_[nm_]...; view content gms3_vc_mme_vin_dtl_<loc>",
     "MME only", "mme/dao/MMEDocumentManagementDAO (49 + 6 sites); mme/dao/MMEDocumentBatchDAO.tables(locale, modelType)",
     "Job log: no 'Table ... doesn't exist'; rows written in gms3_dmt_en_uk_nm_imdoc for an en-UK NewM job", None, None],
    ["M3", "MME-only columns", "ESI row DC_STEERING_TYPE_INFO; VIN row DC_VIN_WMI_CODE; master VIN table VINMAST (MC: VINMASTER)",
     "Carried by the batch inserts", "MME only", "MMEDocumentBatchDAO.saveNewDocuments / updateDocuments",
     "Check DC_VIN_WMI_CODE filled in gms3_dmt_en_uk_nm_vin after the CX-5 job", None, None],
    ["M4", "Model types", "NewM / OldM only (MC Dev schedules checked); no OEM folders",
     "OEM handling of MC (E4 / P2: Isuzu, Nissan, Suzuki) not ported - not applicable for MME", "N/A for MME", "-",
     "-", None, "E4 / P2 do not apply to MME."],
    # ---- M2. Kapture ----
    ["M5", "User groups / view", "addUserGroups(USERGROUPS_MME_*) TECHNICIAN, AFTERMARKET, CORPORATE; addViews(VIEW_MME)",
     "Kapture reference keys TECHNICIAN_CONTENT_KEY, AFTERMARKET_CONTENT_KEY, CORPORATE_CONTENT_KEY; view MME_KEY",
     "Same pattern as MC (own keys)", "kapture.usergroups.mme.refkeys / kapture.views.mme.refkeys (application.properties, application_mcdev.properties); KaptureContentService(groups, views, keepLink)",
     "Kapture document: user groups and view as listed. The 4 keys must exist in Kapture on MC Dev.", None, "Keys to be confirmed with Kapture before the first MME run."],
    ["M6", "Create / modify", "IQServiceClient createContent / modifyContent one document at a time", "processDocumentsInBatches + Kapture REST (create / bulk update), as MC; documents grouped per locale + model type",
     "Yes", "StartMMEConversionImpl.processDocumentsInBatches / applyKaptureResult; MMEDocumentBatchDAO.findExistingDocuments",
     "Report 'KAPTURE OPERATION' (was INFO MANAGER OPERATION): created / updated counts match the source", None, None],
    ["M7", "Categories", "getCategoryByReferenceKey per category (InfoManager)",
     "One batched lookup in k_categories (KaptureLookupDAO) - incl. MME VIN categories (carline + WMI ___ + VDS + VIS) and WD steering LHD / RHD categories",
     "Yes", "StartMMEConversionImpl.locateCategoriesInKapture; CategoryUtils.addVINCategoriesForMME; readESICategoryTextFileForMME_WD",
     "Category report: no NOT FOUND for the CX-3 / CX-5 esicat categories. A NOT FOUND -> missing-carline-categories (Job 8).", None, None],
    ["M8", "Publish", "Document status from the schedule screen (could create Published)",
     "Always Draft; the Publish Content job publishes (key MGSS_PUB_MME_) - status choice hidden on the screen for MME",
     "Yes", "Schedule.scheduleMMEOperation (flag.value.draft); schedule.jsp; publish/PublishMarket (MC / MME); PublishDAO; schedule.name.publish.mme.key",
     "1) Kapture documents Draft after the job 2) Publish Content from History: documents Published, vc_content_status Published in gms3_vc_mme_vin_dtl_<loc>", None, None],
    ["M9", "Items being published", "-", "A schedule for an item (locale + model + manual type) that a pending / running publish job holds is refused (MC and MME publish keys)",
     "Yes", "Schedule.scheduleMMEOperation / scheduleMCOperation; PublishDAO.runningItems", "Schedule the same item while its publish job runs: error message", None, None],
    ["M10", "Delete", "Delete.txt + InfoManager UNPUBLISH category",
     "No delete file: a document ACTIVE in gms3_dmt_<loc>_[nm_]imdoc for the material folder that the ESI txt no longer lists is unpublished in Kapture (kapture.unpublish)",
     "Yes (user decision)", "WithdrawnDocumentsFinder.find(..., mme=true) -> MMEDocumentManagementDAO.getActiveDocumentsForMaterialFolder; StartMMEConversionImpl.deleteContent",
     "Remove one line from esicat.txt, run again: schedule screen shows 1 for deletion, document unpublished, imdoc row inactive", None, None],
    ["M11", "Master Data with Content", "-", "Load type radio for MME too; the job loads the master data before the content (MasterDataLoad), its reports in the zip",
     "Yes (user decision)", "Schedule.java (masterDataWithContent MC or MME); schedule.jsp; StartMMEConversionImpl (MasterDataLoad.run / writeReports)",
     "Schedule with 'Master data with content': master data reports in the job zip; a master data failure fails the job", None, None],
    # ---- M3. content ----
    ["M12", "ENT stylesheet", "addStyleLinkForEntDocumentsForMMEMarket: MME css link appended, other links kept",
     "Every <link> removed first, then /library/MAZDA/GMS3_CUSTOM/mazda-css-mme.css appended (rule O1, as MC A6)",
     "Same rule, MME css", "ConversionUtils.addStyleLinkForEntDocumentsForMMEMarket / entStyleHrefMME; GMS3_MME_CUSTOM_CONTENT_CSS",
     "CX-3 / CX-5 ENT documents: exactly one link, mazda-css-mme.css", None, None],
    ["M13", "Kept stylesheet links (sanitizer + preview)", "-",
     "The sanitizer keeps the MME ENT css and content.css / contents.css (MC: the MC css)", "Same rule, MME css",
     "ConversionUtils.isConversionStyleLinkMME; KaptureContentService 3-arg constructor; PreviewBuilder(scheduleId, true)",
     "Kapture document source and preview page carry the same link", None, None],
    ["M14", "RDF / XML documents", "ENT folder .xml -> processARDFFile (XML_DOCUMENT); XML copied to OKAssets",
     "OASIS_FILE_LOCATION = web path of the xml (setOasisDirectoryPath), SM channel; copied by processENTDirectoryForOkAssets",
     "MME only", "StartMMEConversionImpl.processARDFFile; KaptureArticleBuilder (OASIS_FILE_LOCATION)",
     "CX-3: 2015_CX3_OASIS.xml -> document with OASIS_FILE_LOCATION, file present in OKAssets", None, None],
    ["M15", "WD voltage map links", "LINK_URL_MME_WINDOW_JS", "Same; the preview uses LINK_URL_MME_WINDOW_JS for MME jobs",
     "Same rule, MME key", "PreviewBuilder(scheduleId, true)", "WD preview of an MME job: voltage map links open", None, None],
    # ---- M4. text files / MME-only steps ----
    ["M16", "vin.txt", "readVINTextFileForMME: 9 tokens (adds WMI); carline names from gms3_mdm_vin_detail at en_UK",
     "Unchanged result (no InfoManager in it); MDM read at en_UK, no fallback. Carline names now read in bulk (O12): one SELECT per chunk of carline codes, matched in memory - was one SELECT per VIN line (shared with MC vin.txt, ja-JP).", "MME only", "ConversionUtils.readVINTextFileForMME / readSCMVINTextFileForMME; MMEDocumentManagementDAO.getModelDetails",
     "VIN rows with WMI; blank carline name = VIN missing in MDM (expected, no fallback). Job log: 'getModelDetails :: n lines, ... matched ... ms' once per file", None, "Open point O12 (CLOSED)."],
    ["M17", "scmvin*.txt (WM only)", "readSCMVINTextFileForMME + startSCMVINProcessing[ForFaceLift]: 9th token = source file -> document id -> gms3_dmt_schm_mapping",
     "Same result, in bulk: document ids and existing gms3_dmt_schm_mapping rows read per chunk of paths (IN list), one UPDATE (CASE per column) and one INSERT batch per chunk. Was up to 4 statements per line, each on its own connection. Same SCMV001-003 statuses; the same key twice in the file = one row holding the last line (as insert-then-update).",
     "MME only", "ConversionUtils.readSCMVINTextFileForMME; StartMMEConversionImpl.startSCMVINProcessing; MMEDocumentManagementDAO.saveSCMVinProcessingDetails; printSCMVINMappingReport",
     "CX-5: SCM VIN report lists the 165 rows; all point to id000000801200.ent, which is NOT in the folder -> SCMV003 (expected). Job log: one 'saveSCMVinProcessingDetails :: ... ms' line", None, "Open point O9 (CLOSED)."],
    ["M18", "VIN -> manual type mapping", "startVINManualTypeMapping: gms3_mdm_vin_ml_mapping (VIN_ML_WD / WM ...)",
     "Same result, in bulk: existing ACTIVE rows of the locale read per chunk of carlines, one UPDATE ... WHERE VIN_ML_ID IN (...) per chunk (same values for all), one INSERT batch per chunk; a VIN listed twice (vin.txt + master VINs) inserted once. Was SELECT + UPDATE / INSERT per VIN.",
     "MME only", "StartMMEConversionImpl.startVINManualTypeMapping; MMEDocumentManagementDAO.saveVinManualTypeMapping / writeVinManualTypeMapping; printVinMLMappingReport",
     "VIN ML mapping report in the zip; gms3_mdm_vin_ml_mapping VIN_ML_WM_MAPPING set for the CX-3 / CX-5 VINs, no duplicate rows. Job log: 'saveVinManualTypeMapping :: n VINs, u updated, i inserted ... ms'", None, "Open point O9 (CLOSED)."],
    ["M19", "Display order / CD text files", "Keys and checks on Level*CODE",
     "The files carry NAMES - keys and checks on Level*NAME (as MC); name validation blocks removed. Display order staged in gms3_dmt_el_gr_nm_dispord and cleaned at the end (existing pattern, user decision). A document's display order rows are read from the same rows held in memory (index by path) - no SELECT per document.",
     "Yes", "StartMMEConversionImpl.returnKey / indexDisplayOrder / applicableDisplayOrderForDocument; MMEDocumentManagementDAO.saveCDProcessingDetails",
     "Display order rows in gms3_dmt_en_uk_nm_dispord with the d01-d07 names, same as a document's display order lines; no duplicates collapsed", None, "Open point O8 (CLOSED)."],
    # ---- M5. job mechanics ----
    ["M20", "Counters", "Data Preparation / OKAssets counts one UPDATE per document; failure helper raised the processing count",
     "Buffered (flushed per phase); failures raise the failure count", "Yes",
     "bufferDataPreparationCount / flushDataPreparationCount; OkAssetsCountBuffer.begin / end", "Schedule screen counts match the reports", None, None],
    ["M21", "No batched UPDATE / DELETE (MySQL Router)", "updateInnerLinkMappingStatus batched UPDATE; batched DELETEs",
     "One UPDATE per chunk (CASE + IN list); dead batched DELETEs removed; child-table deletes by IN list",
     "Yes - MC now the same (O10)", "MMEDocumentManagementDAO / MCDocumentManagementDAO.updateInnerLinkMappingStatus; MMEDocumentBatchDAO.deleteViewContent", "Inner link pass finishes without Router errors", None, None],
    ["M22", "Inner links", "One pass, InfoManager modify per document", "Two passes: documents whose links map only after their targets exist are updated again (as MC B1)",
     "Yes", "StartMMEConversionImpl (inner links); MMEDocumentBatchDAO.updateDocumentsForInnerLinks", "CX-3 ENT inner links open the target document in Kapture and preview", None, None],
    ["M23", "Preview", "-", "Offline preview DMT_PREVIEW/<job>/ as MC (view content of the locale)", "Yes",
     "PreviewBuilder(scheduleId, true); preview.finish(getViewContentDataList()); servlet/History, Preview (thread prefix per market)",
     "History: Preview link for the MME job; pages = Kapture content", None, None],
    ["M24", "E-mail", "FetchUserProfileImpl (InfoManager)", "UserProfileDAO.getUserEmail", "Yes", "StartMMEConversionImpl", "Notification e-mail received", None, None],
    ["M25", "View content status", "-", "VC_CONTENT_STATUS = Draft on insert; Published by the publish job",
     "Yes", "MMEDocumentManagementDAO (view.content.status.draft); db/dmtgt_mme_pre_deployment.sql (column on the 14 tables)",
     "Run the SQL script BEFORE deploying; check query returns 14 rows", None, None],
    ["M26", "Master data categories (Master Data with Content)", "Master Data Sync screen -> InfoManager categories",
     "Same hierarchy, created through the Kapture category API (parent first). MC + MME: CARLINE sheet -> under CARLINE (carline, level 1); ALL_VIN / VIN sheet -> under the same tree: CARLINE > carline > WMI > VDS > VIS start > VIS range (levels 2-5). MNAO (not ported yet): CARLINE sheet -> MODEL_YEAR > model > model + year; ALL_VIN / VIN sheet -> VIN > WMI > WMI+VDS > VIS range.",
     "Yes (shared code)", "masterdata/KaptureCategorySync; autosync/dao/CarlineDataDAO (MC / MME), ModelYearDataDAO + VINRangeDataDAO (MNAO only); PARENT_REF_KEYS_CARLINE / _MODEL_YEAR / _VIN_RANGE",
     "MASTER_DATA_CATEGORY_REPORT: each level CREATED / NO CHANGE under the expected parent", None, None],
]

# Results of the runs (SC Dev from the workstation, then MC Dev) - Result / Remarks of the MME checklist
MME_RESULTS = {
    "M2": "SC Dev 5438-5443, MC Dev 5603: rows in gms3_dmt_en_uk_nm_* tables, no 'doesn't exist' errors.",
    "M5": "Keys confirmed on MC Dev (user, 2026-10-05); MC Dev 5603: 470 documents created with the MME groups / view.",
    "M6": "SC Dev 5438 created, 5439 / 5443 update path; MC Dev 5603: 470/470 created (100 per request, ~0.6-0.8 s per document).",
    "M7": "SC Dev 5438: missing categories listed (T8); MC Dev 5603: 0 NOT FOUND.",
    "M8": "MC Dev publish 5610 (from 5603): 470/470 published in 10 requests, 0 failed; preview folder removed.",
    "M10": "SC Dev 5440: SM3465162 (CX-3) and SM3465175 (CX-5) unpublished, rows removed, preview rows DELETED by 5440.",
    "M11": "SC Dev 5438-5443 loaded master data first. MC Dev 5603: the 7 MDM Excel files 'Permission denied' for Tomcat in /s3kapture/sourcecontent/Test/MDM -> job Failure as designed (infra: grant read access).",
    "M15": "MC Dev 5617: voltage map JS links looked up in the MME tables with LINK_URL_MME_WINDOW_JS (see T11).",
    "M16": "Job logs: 'getModelDetails :: ... matched' once per file, 8-20 ms on MC Dev.",
    "M17": "SC Dev 5438: 165 x SCMV003 (expected, target document not in the folder).",
    "M18": "MC Dev 5603: 'saveVinManualTypeMapping' 3 + 1 + 1 rows inserted, ~245 ms each.",
    "M19": "MC Dev 5603: 5,214 + 461 display order rows staged in < 1 s; display order + CD 470/470. Repeated names in the tree = several display order lines (accepted).",
    "M21": "MC Dev 5603: no Router errors.",
    "M23": "SC Dev 5438-5443 and MC Dev 5603: preview built (470 documents, 5 VINs). PRE-DELIVERY INSPECTION has no placement: its d01.txt line has no level name (data).",
    "M24": "MC Dev 5603 and publish 5610: e-mail sent.",
    "M25": "MC Dev: pre-deployment SQL run before the war (2026-10-05); publish 5610 set Published.",
    "M26": "SC Dev 5438: categories created under CARLINE / ESI; 5 ESI rows answered 409 'name already exists under the same parent' (Kapture data). MC Dev 5617: 337/343 OK, 1 create time-out + 5 ESI 409 (O14). Sync status now set in bulk (O13). MNAO: to be ported with the MNAO market.",
}
for _r in MME_ROWS:
    if _r[0] in MME_RESULTS:
        _r[7] = "OK"
        _r[8] = MME_RESULTS[_r[0]] if not _r[8] or "to be confirmed" in _r[8] else _r[8] + " " + MME_RESULTS[_r[0]]

MME_TEST_HEAD = ["Ref", "Content (MC Dev / D:\\Kapture_WD\\SourceContent\\GMS3_Content\\MME\\en-UK)", "What it covers", "Expected", "Result", "Remarks"]
MME_TEST = [
    ["T0", "Before the first run", "DB script + properties",
     "db/dmtgt_mme_pre_deployment.sql run on kapture_dc (check query: 14 rows); application_mcdev.properties swapped in; the 4 Kapture keys of M5 exist", None, None],
    ["T1", "NewM_CX-3_dk / WM / FL0000 (E-WEE-03K1_XML: 7 ENT + 2015_CX3_OASIS.xml; d01-d05, CD00, vin.txt)", "M6, M7, M12-M14, M16, M19, M22, M23",
     "7 ENT documents + 1 RDF document created as Draft; mazda-css-mme.css the only link; OASIS_FILE_LOCATION set; display order + CD written; preview available", None, None],
    ["T2", "NewM_CX-5_kf / WM / FL0000 (E-WEE-72A1_XML; d01-d07, CD00, vin.txt, scmvin.txt)", "M3, M16-M18",
     "Documents created; SCM VIN report: 165 rows not resolved (all point to id000000801200.ent, not in the folder) - expected; VIN ML mapping report", None, None],
    ["T3", "Run T1 again unchanged", "M6 (update path)", "All documents updated, none created", None, None],
    ["T4", "T1 with one line removed from the ESI txt", "M10", "Schedule screen: 1 for deletion; document unpublished in Kapture; imdoc row inactive", None, None],
    ["T5", "Publish Content for the T1 job (History)", "M8, M9, M25", "Documents Published; vc_content_status Published; scheduling CX-3 WM while the publish job runs is refused", None, None],
    ["T6", "T1 with load type 'Master data with content'", "M11", "Master data loaded first, its reports in the zip", None, None],
    ["T7", "NewM_MAZDA2_Hybrid_KBAC3 / SH / FL0003 (461 documents)", "Volume", "1 expected failure: d01 lists E-TZE-6HB3_PDF/Introduction.pdf, the PDF folder is absent", None, "BT-50 WM (5,214 documents, ~3.3 h) only after T1-T7."],
    ["T8", "T1 / T2 on SC Dev - categories NOT pre-checked in k_categories (user decision 2026-10-05)", "M7 - missing categories path (never occurred for MC)",
     "Any ESI / carline-VIN / steering category not in SC Dev k_categories is listed NOT FOUND in the category report; check how the documents carrying it are reported (not silently loaded without it) and that the job completes. Feed the report to missing-carline-categories (Job 8), rerun: categories found.", None,
     "First run of this path for any market."],
    ["T9", "NewM_BT-50_tf / WM / FL0012 (E-WZE-8FC2_PDF: 9 PDFs, up to 86 MB)", "PDF documents, large attachments",
     "9 PDF documents created with their attachment", None, None],
    ["T10", "MC Dev end to end: BT-50 PDF + MAZDA2 Hybrid SH HTML + CX-5 km PDF, then Publish Content", "M2-M25 on MC Dev (MySQL Router, Linux paths)",
     "All documents created as Draft, then published; no Router errors; e-mail sent", None, None],
    ["T11", "An MME WM HTML folder with wiring diagrams", "M15 - window.js / voltage map JS",
     "window.js and the voltage map JS generated for MME (LINK_URL_MME_WINDOW_JS); WD preview voltage links open", None, None],
    ["T12", "MC Dev re-run with master data readable + WD content (5617: BT-50 PDF, Hybrid SH, CX-5 km e-WD + WM PDF)", "M11, M26, update of published documents, volume",
     "Master data loaded and synced to Kapture; published documents updated as new drafts; WD created", None, None],
]

MME_TEST_RESULTS = {
    "T0": ("OK", "MC Dev 2026-10-05: pre-deployment SQL run, war deployed; 4 Kapture keys present."),
    "T1": ("OK", "SC Dev 5438 (CX-3 + CX-5: 12 documents created)."),
    "T2": ("OK", "SC Dev 5438: 165 x SCMV003 as expected; VIN ML mapping in bulk."),
    "T3": ("OK", "SC Dev 5439: updates only, ~2 min (master data already in Kapture)."),
    "T4": ("OK", "SC Dev 5440: CX-3 id0102q2004400.ent and CX-5 id000000000100.ent unpublished, rows removed."),
    "T5": ("OK", "MC Dev publish 5610 for job 5603 (instead of the T1 job): 470/470 published."),
    "T6": ("OK", "SC Dev 5438-5443. MC Dev 5603: MDM files not readable (Permission denied) -> job Failure as designed; infra to grant read access."),
    "T7": ("OK", "SC Dev 5441: 459/460 (1 Kapture ARTICLE_CREATION_UNVERIFIED, fixed by Kapture, created in 5443 as SM3465826). 5443: 3 transient Kapture 'write conflict' rollbacks. MC Dev 5603: 460/460."),
    "T8": ("OK", "SC Dev 5438: 348 missing categories reported, job completed (Failure). MC Dev 5603: 0 missing."),
    "T9": ("OK", "SC Dev 5441: 9 x ADD_ARTICLE_ERROR (Kapture 500) -> fixed by Kapture; 5442: 9/9. MC Dev 5603: 9/9."),
    "T10": ("OK", "MC Dev 5603 (6 min 42 s): 470/470 created, all folders SUCCESS, 0 missing categories; job Failure only from the MDM file permission (T6). Publish 5610 (6 min 18 s): 470/470."),
    "T11": ("OK", "MC Dev 5617: NewM_CX-5_km e-WD E-DEE-69P1 - 1,851 HTML5 pages created, voltageMAP.js / voltagelinkMAP.js copied and their links looked up in gms3_dmt_en_uk_nm_imdoc with the MME link (LINK_URL_MME_WINDOW_JS). 80 + 4 links not resolved: they point to WM XML documents (NewM_CX-5_km/WM/FL0000/E-WZE-69P1_XML) not loaded yet - re-check after loading that folder. No window.js in this content."),
    "T12": ("OK", "MC Dev 5617 (44 min, 2,321 documents; master data 18 min of it): master data read (264 added, 20 updated); 337/343 categories OK - 1 carline VDS create timed out (120 s), 5 ESI 409 (Kapture data). Content: 2,311 OK; 10 Hybrid updates rolled back by Kapture ('Plugin instructed the server to rollback the current transaction' = DB write conflict, Kapture side; re-run updates them)."),
}
for _t in MME_TEST:
    if _t[0] in MME_TEST_RESULTS:
        _t[4], _rem = MME_TEST_RESULTS[_t[0]]
        _t[5] = _rem if not _t[5] else _t[5] + " " + _rem

# ---------------------------------------------------------------------------------------------
# MNAO - the MC / MME flow for en-US, es-MX, fr-CA, en-CA (audit tmp/DMT_MNAO_Audit_20261006.md rev 4)
# ---------------------------------------------------------------------------------------------
MNAO_HEAD = ["Ref", "Area", "Legacy MNAO (DMT_Optimized, InfoManager)", "DMTKapture MNAO (now)", "Same as MC / MME?",
             "Code / config", "How to validate", "Result", "Remarks"]

MNAO_ROWS = [
    ["N1", "Locales", "en-US, es-MX, fr-CA, en-CA; each locale an InfoManager document",
     "Each locale its own Kapture master article (as MME, Q7); Kapture locale from the property <xx-yy> (en-us=en_US ...)",
     "Yes (as MME)", "StartConversionMNAOImpl (kaptureLocale = getProperty(locale.toLowerCase()))",
     "Kapture: no translation parent; locale = the schedule locale", None, None],
    ["N2", "Folder layout", "<locale>\\<model>\\<material> (no manual type, no facelift folder)",
     "MC / MME layout <locale>\\<model>\\<manual>\\<FLxxxx>\\<material>; the model folder has no model type (CX-5_KM) and is parsed with an empty model type in front",
     "Yes", "servlet/Schedule identifyFileItemsOperationForMC; mc/utils/ConversionUtils.modelFolderForParsing",
     "Schedule screen lists CX-5_KM / WM / FL0000 / E-WUE-69P1_MY2026_2026_XML with model CX-5", None, None],
    ["N3", "Existing DMT rows", "Paths without a facelift folder",
     "db/dmtgt_mnao_pre_deployment.sql: dc_facelift_folder_name = FL0000 and 'FL0000\\' inserted after the 3rd backslash of every MNAO path (imdoc, cat, vin, vinattr, vinmaster, cvc, inrlks, both view content tables)",
     "MNAO only", "db/dmtgt_mnao_pre_deployment.sql (section 8 = checks)", "Section 8 of the script: every count 0", None,
     "Local kapture_dc 2026-10-06: 15m28s, section 8 all 0. OKAssets folders are restructured manually on the server (R2)."],
    ["N4", "VIN files", "vin.ent + vin_attribute.ent",
     "vin.txt, vin1.txt ... in the facelift folder (file.vin.ent + .txt, not scmvin); a line applies when manual + FL + material match; 11 fields incl. model and year (a line with fewer fields is a VIN file issue)",
     "Yes (MNAO reader readVINTextFileForMNAO)", "mc/utils/ConversionUtils.readVINTextFileForMNAO",
     "vin.txt of BT-50 WM FL0012 (2 lines) read; a 10-field line is reported, not loaded", None, "vin.ent / vin_attribute.ent are no longer read (user 2026-10-06)."],
    ["N5", "Display order files", "-",
     "d01.txt, d02.txt ... (file.name.displayorder, not delete); a document takes the display order file its vin.txt line names - matched on the complete file name, exactly as MC / MME",
     "Yes", "same comparison as StartMCConversionImpl / StartMMEConversionImpl (getDisplayOrderSourceFileName equals getDisplayOrderTextFileName)",
     "Lines naming d01.txt and d02.txt map their documents to those files", None, None],
    ["N6", "ESI file", "esicat.txt (MNAO layout)", "esicat.txt read with the plain reader (8 fields, no steering column)",
     "Partly (no steering column)", "ConversionUtils.readESICategoryTextFile", "Documents carry their ESI categories", None, None],
    ["N7", "Steering", "-", "rhdlhdindicator META of the source html", "MNAO only", "StartConversionMNAOImpl.readSteeringFromMeta",
     "Steering category on an html document with the META", None, None],
    ["N8", "Categories", "getCategoryByReferenceKey per category (InfoManager)",
     "MNAO reference keys (conversion/impl/CategoryUtils) fed from the vin.txt lines: manual type, MODEL_YEAR <MODEL>_<YEAR>, ESI, steering, CVC, VIN (WMI+VDS+range), master VIN, engine / mission; drive axle / body type from d01 as MC / MME. Missing category = Missing Category report, not created (Q8)",
     "Same flow, own keys", "mc/utils/CategoryUtils.performCategoriesOperationForMNAO",
     "Category report: no NOT FOUND for the sample once Master Data with Content ran", None, None],
    ["N9", "User groups / view", "USERGROUPS_MNAO_* / VIEW_MNAO",
     "kapture.usergroups.mnao.refkeys = TECHNICIAN_CONTENT_KEY, AFTERMARKET_CONTENT_KEY; kapture.views.mnao.refkeys = MNAO_KEY",
     "Same pattern", "application.properties, application_mcdev.properties", "Kapture document: user groups and view as listed", None,
     "MNAO_KEY not verified in Kapture yet."],
    ["N10", "Create / modify", "IQServiceClient one document at a time",
     "Kapture REST create / bulk update, always unpublished (Draft); DMT tables through MNAODocumentBatchDAO (batched INSERTs, one UPDATE per row)",
     "Yes", "conversion/dao/MNAODocumentBatchDAO", "Report KAPTURE OPERATION: created / updated counts match the source", None,
     "Legacy bug fixed: a document with more than 100 VINs wrote its first 100 again."],
    ["N11", "Publish", "Published / Draft radio", "Draft load + Publish Content job MGSS_PUB_MNAO_<date>_<job> from the Preview page; no status choice on the schedule screen",
     "Yes", "publish/PublishMarket.MNAO (DC_IM_DOC_FLAG = active flag, DC_IM_DOC_STATUS = published flag, both view content tables); schedule.name.publish.mnao.key",
     "After publish: imdoc DC_IM_DOC_STATUS = 4, vc_content_status Published in gms3_vc_model_year_details and gms3_vc_vin_details", None, None],
    ["N12", "View content", "processViewContentData after the documents (VIN ent rows)",
     "Written in the display order phase: model year rows (MODEL_YEAR category x display order row) and VIN rows (vin.txt lines + master VIN rows, en-CA year ranges skipped), model and year on every VIN row, display order names 1-6 + sequence no, vc_content_status Draft; one DELETE IN + INSERT batch per chunk",
     "Same pattern (MNAO tables)", "MNAODocumentManagementDAO.saveDisplayOrderDetails",
     "gms3_vc_vin_details rows of the job: vc_vin_model / vc_vin_year filled, vc_content_status Draft", None, None],
    ["N13", "Display order staging", "-", "Shared temp table gms3_dmt_el_gr_nm_dispord, as MME; the per-locale gms3_dmt_<loc>_dispord is only cleaned",
     "Yes (as MME)", "ScheduleDAO", "Temp rows of the job removed at the end", None, None],
    ["N14", "Scheduling validations", "-",
     "As MC / MME: no display order and no VIN; VIN file issues; display order but no VIN; VIN but no display order; every display order file named by vin.txt exists",
     "Yes", "servlet/Schedule (VIN / display order block)", "A folder with vin.txt naming a missing d02.txt is refused on the schedule screen", None, None],
    ["N15", "Deleted documents", "Delete.txt -> unpublish category",
     "WithdrawnDocumentsFinder(market MNAO): active rows of the material folder not in esicat.txt -> Kapture unpublish",
     "Yes", "Schedule (count for deletion); StartConversionMNAOImpl", "Schedule screen: n for deletion; documents unpublished", None, None],
    ["N16", "OKAssets", "conversion/utils/XcopyUtil, <loc>/<model>/<material>",
     "mc/utils/XcopyUtil: /library/MAZDA/<CH>/<loc>/<model>/<manual>/<fl>/<material>/ (lower case); window.js OKAssets path with the facelift folder",
     "Yes", "mc/utils/XcopyUtil", "Images / html5 load in the preview from the new path", None,
     "Existing window.js files are corrected manually (user)."],
    ["N17", "Voltage map JS", "Links written with the legacy MNAO rule",
     "Links matched as authored in the content - no facelift folder inserted; link property LINK_URL", "Same rule as MC",
     "WiringDiagramUtils.performVoltageMap(Link)JSOperationForMNAO", "e-WD html5: voltage map opens the target document", None,
     "The links of the current sample are wrong; the real content carries the facelift folder (user)."],
    ["N18", "Stylesheet links", "ENT: mazda.css (GMS3_CUSTOM_CONTENT_CSS) appended",
     "ENT: every link removed, /library/MAZDA/GMS3_CUSTOM/mazda.css (GMS3_CUSTOM_CONTENT_CSS) the only link; HTML: content.css / contents.css of the html folder; the sanitizer and the preview keep exactly these",
     "Same rule, MNAO css", "ConversionUtils.addStyleLinkForEntDocumentsForMNAOMarket / isConversionStyleLinkMNAO; PreviewBuilder(scheduleId, market)",
     "Kapture source of an ENT document: one link, mazda.css", None, None],
    ["N19", "Master Data with Content", "-",
     "Radio on the schedule screen; CARLINE MNAO layout (YEAR START / YEAR END -> mdm_crln_year_start / end, key code + name + WMI + year start); MODEL_YEAR and VIN categories sent to Kapture; DRIVELINE_AXLE_TYPE / BODY_TYPE files (all markets)",
     "Yes (shared code)", "masterdata/MdmExcelLoader, masterdata/KaptureCategorySync",
     "MASTER_DATA_LOAD_REPORT rows for US_EN-US_CARLINE.xlsx; category report rows for MODEL_YEAR / VIN", None,
     "Drive axle / body type also land in MC and MME."],
    ["N20", "Inner links", "One UPDATE per document", "One CASE / IN UPDATE per chunk", "Yes", "MNAODocumentManagementDAO.updateInnerLinkMappingStatus",
     "No Router error; links mapped", None, None],
    ["N21", "Master VIN flag", "if(addToVINMaster=true) - always true", "if(addToVINMaster==true) (R3)", "-", "MNAODocumentManagementDAO", "Master VIN rows only for the pseudo models", None, None],
    ["N22", "MME-only phases", "-", "No CD data, SCM VIN or VIN - manual type mapping for MNAO", "N/A for MNAO", "-", "-", None, None],
    ["N23", "Model year ranges", "One SELECT per document", "Read once per job (cache cleared at job start)", "-", "MNAODocumentManagementDAO.getModelYearRangesList", "-", None, None],
]

MNAO_TEST_HEAD = ["Ref", "Content (D:\\Kapture_WD\\SourceContent\\GMS3_Content\\MNAO\\en-US)", "What it covers", "Expected", "Result", "Remarks"]
MNAO_TEST = [
    ["U0", "Before the first run", "DB script + properties",
     "db/dmtgt_mnao_pre_deployment.sql run (section 8 all 0); application_mcdev.properties swapped in; the keys of N9 exist in Kapture", None, None],
    ["U1", "BT-50_TF / WM / FL0012 / E-WZE-8FC2_MY2026_HTML (html 12,006 + pdf 3; d01 5,214; vin.txt 2 lines)", "N2, N4, N5, N8, N10, N12",
     "Documents created as Draft; view content rows with model CX90 / CX90 2ROW and year; preview available", None, "Volume run - after U2 / U3."],
    ["U2", "CX-5_KM / WM / FL0000 / E-WUE-69P1_MY2026_2026_XML (ent.gms3 4,185)", "N4 (10-field line), N18",
     "vin.txt line without model reported as a VIN file issue on the schedule screen until business adds the model (Q-A); then ENT documents with mazda.css only", None, None],
    ["U3", "CX-5_KM / e-WD / FL0000 / E-DUE-69P1 (html5 920)", "N16, N17",
     "WD documents; OKAssets under /library/MAZDA/WIRING_DIAGRAMS/en_us/cx-5_km/e-wd/fl0000/e-due-69p1/", None, None],
    ["U4", "Run U3 again unchanged", "N10 (update path)", "All documents updated, none created", None, None],
    ["U5", "U3 with one line removed from esicat.txt", "N15", "1 for deletion; document unpublished; imdoc row inactive", None, None],
    ["U6", "Publish Content for the U3 job", "N11", "Published; vc_content_status Published in both view content tables", None, None],
    ["U7", "U3 with 'Master data with content' and US_EN-US_*.xlsx (+ DRIVELINE_AXLE_TYPE / BODY_TYPE files)", "N19",
     "Master data reports in the zip; carline years loaded; MODEL_YEAR / VIN categories in Kapture", None, None],
    ["U8", "vin.txt lines naming d01.txt and d02.txt", "N5, N14", "Each document takes the file its line names; a vin.txt naming a missing d02.txt refuses the schedule", None, None],
]

# ---------------------------------------------------------------------------------------------
# Sheet 4 - legend
# ---------------------------------------------------------------------------------------------
LEGEND = [
    ["Placeholder / key", "Meaning / MC Dev value"],
    ["<loc>", "Locale folder in OKAssets, lower case: ja_jp (source folder ja-JP; property ja-jp=ja_JP)"],
    ["<CH>", "SERVICE_MANUALS | OTHER_SERVICE_MANUALS | WIRING_DIAGRAMS (channel folder upper case, spaces -> _)"],
    ["<model> <mt> <fl> <mat>", "Model folder, manual type, facelift folder, material folder - lower case"],
    ["SERVER_LIBRARY_DIRECTORY / REPOSITORY", "library/ and MAZDA -> every path starts /library/MAZDA/"],
    ["SERVER_OKASSETS_PHYSICAL_PATH", "/s3kapture/ROOT/resources/sites/KAPTURE/content/ (MC Dev)"],
    ["okassets.web.context", "/content - the browser reaches OKAssets at /content/library/...; used by the preview, reports dialog, e-mails - NOT inside Kapture documents"],
    ["LINK_URL / LINK_URL_MC_WINDOW_JS", "content/ (legacy: index?page=content&id= / index?page=detail_mc&id=)"],
    ["GMS3_INFOCENTER_APPLICATION_CONTENT", "mazdagms3/ (legacy: dokinfoctr/mazdagms3/) - stays, user decision"],
    ["OTHER_MANUAL_LINK_RULES", "result_mme, result_mc, result (B3)"],
    ["kapture.sanitize.strip.tags", "meta,script,iframe,object,embed,link"],
    ["preview.physical.path / preview.web.path", "<OKAssets>/library/MAZDA/GMS3_CUSTOM/DMT_PREVIEW/<schedule id>/ - browser /content/library/MAZDA/GMS3_CUSTOM/DMT_PREVIEW/<schedule id>/"],
    ["DMT_PREVIEW/<job>/", "<document id>.html, files/<document id>/<pdf>, wd/<document id>/<page>, _assets/preview_page.js, preview.json"],
    ["<loc> (MME)", "MME table locale, lower case with _: en_uk (source folder en-UK); 14 locales - M1"],
    ["GMS3_MME_CUSTOM_CONTENT_CSS", "mazda-css-mme.css - the ENT css of MME (M12)"],
    ["GMS3_CUSTOM_CONTENT_CSS", "mazda.css - the ENT css of MNAO (N18)"],
    ["<loc> (MNAO)", "MNAO table locale: en_us, es_mx, fr_ca, en_ca (source folder en-US ...) - N1"],
    ["Result column", "OK / NOT OK / N/A - fill while validating a run (note the schedule id in Remarks)"],
]


def style_sheet(ws, head, rows, widths):
    hfill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = border
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = border
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def highlight(ws, col_letter, words, color):
    fill = PatternFill("solid", fgColor=color)
    for c in ws[col_letter][1:]:
        if c.value and any(str(c.value).startswith(w) for w in words):
            c.fill = fill


def result_dropdown(ws, col_letter, n):
    dv = DataValidation(type="list", formula1='"OK,NOT OK,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (col_letter, col_letter, n + 1))


wb = Workbook()
ws = wb.active
ws.title = "Checklist"
style_sheet(ws, CHECK_HEAD, CHECK, [6, 24, 24, 38, 40, 44, 44, 13, 36, 40, 10, 36])
highlight(ws, "H", ["YES", "Fixed"], "FCE4D6")
highlight(ws, "H", ["YES - OPEN"], "FFF2CC")
highlight(ws, "L", ["OPEN FOR DISCUSSION"], "FFF2CC")
result_dropdown(ws, "K", len(CHECK))

ws2 = wb.create_sheet("Display by document type")
style_sheet(ws2, TYPE_HEAD, TYPES, [6, 22, 32, 50, 40, 50, 18, 40, 10])
result_dropdown(ws2, "I", len(TYPES))

wsp = wb.create_sheet("Legacy cleaning per process")
style_sheet(wsp, PROC_HEAD, PROC, [6, 26, 22, 60, 40, 40, 30, 34, 20])
note = wsp.cell(row=len(PROC) + 3, column=2,
                value="OPEN FOR DISCUSSION (O7): the sanitizer is generic today; whether to follow legacy exactly per process function is for the user to decide.")
note.font = Font(bold=True, color="9C5700")

ws3 = wb.create_sheet("Open points")
style_sheet(ws3, OPEN_HEAD, OPEN + MME_OPEN, [6, 55, 50, 55, 16, 14])
highlight(ws3, "F", ["OPEN", "TO CONFIRM"], "FCE4D6")
highlight(ws3, "F", ["OPEN FOR DISCUSSION"], "FFF2CC")

wsm = wb.create_sheet("MME checklist")
style_sheet(wsm, MME_HEAD, MME_ROWS, [6, 26, 44, 50, 16, 44, 46, 10, 34])
highlight(wsm, "E", ["MME only"], "DDEBF7")
highlight(wsm, "E", ["N/A"], "EDEDED")
result_dropdown(wsm, "H", len(MME_ROWS))

wst = wb.create_sheet("MME test plan")
style_sheet(wst, MME_TEST_HEAD, MME_TEST, [6, 50, 30, 70, 10, 36])
result_dropdown(wst, "E", len(MME_TEST))

wsn = wb.create_sheet("MNAO checklist")
style_sheet(wsn, MNAO_HEAD, MNAO_ROWS, [6, 26, 44, 50, 16, 44, 46, 10, 34])
highlight(wsn, "E", ["MNAO only"], "DDEBF7")
highlight(wsn, "E", ["N/A"], "EDEDED")
result_dropdown(wsn, "H", len(MNAO_ROWS))

wsu = wb.create_sheet("MNAO test plan")
style_sheet(wsu, MNAO_TEST_HEAD, MNAO_TEST, [6, 50, 30, 70, 10, 36])
result_dropdown(wsu, "E", len(MNAO_TEST))

ws4 = wb.create_sheet("Legend")
for i, r in enumerate(LEGEND):
    ws4.append(r)
ws4.column_dimensions["A"].width = 40
ws4.column_dimensions["B"].width = 120
for c in ws4[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
for row in ws4.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("written", os.path.normpath(OUT), "checklist rows", len(CHECK), "types", len(TYPES), "open points", len(OPEN) + len(MME_OPEN),
      "MME rows", len(MME_ROWS), "MME tests", len(MME_TEST),
      "MNAO rows", len(MNAO_ROWS), "MNAO tests", len(MNAO_TEST))
