# -*- coding: utf-8 -*-
"""Builds DMTKapture_Pending_Test_Tracker.xlsx - every DMT test still to run, to be closed together after Automation.
Re-run after editing the rows below:  python build_pending_test_tracker.py
"""
import io, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "DMTKapture_Pending_Test_Tracker.xlsx")
STATUSES = ["Pending", "In Progress", "Passed", "Failed", "Blocked", "Not Applicable"]

HEAD = ["ID", "Area", "Market", "Test", "Why it is pending", "Steps", "Expected result", "Environment",
        "Depends on", "Priority", "Status", "Job ID / Evidence", "Tested on", "Remarks"]

# ID, Area, Market, Test, Why, Steps, Expected, Env, Depends on, Priority, Status
TESTS = [
    # ---- A. Master data / categories (changes of 2026-10-10 / 11)
    ("T01", "Master Data with Content", "MNAO", "en-CA companion locale - files present",
     "New 2026-10-11: an en-US schedule also processes en-CA (masterdata.companion.locales). Built, not yet run.",
     "1. Put CA_EN-CA_<type>.xlsx files next to the en-US files in the MDM folder.\n2. Schedule MNAO en-US with Master Data with Content.",
     "MASTER_DATA_LOAD_REPORT lists en-CA rows; MASTER_DATA_CATEGORY_REPORT lists en_CA categories; en-CA MDM rows end with sync status Y; "
     "log shows 'companion locale en-CA processed too'.",
     "MC Dev", "MDM country CA / manual language en-CA active", "High", "Pending"),
    ("T02", "Master Data with Content", "MNAO", "en-CA companion locale - no en-CA file",
     "New 2026-10-11: a missing en-CA file must not fail the job.",
     "Schedule MNAO en-US with Master Data with Content, MDM folder WITHOUT CA_EN-CA files.",
     "No 'No master data file found' row for en-CA; log line 'companion locale en-CA :: no master data file'; en-CA rows already at N are sent to Kapture.",
     "MC Dev", "-", "High", "Pending"),
    ("T03", "Schedule - Only content check", "MNAO", "Only content warning counts en-CA",
     "New 2026-10-11: the check now covers the companion locale.",
     "1. Leave at least one en-CA MDM row at sync status N.\n2. Schedule MNAO en-US with Only content.",
     "Warning appears and names 'locale en-US, en-CA'. With no N rows in en-US / en-CA: no warning.",
     "MC Dev", "-", "Medium", "Pending"),
    ("T04", "Conversion re-test", "MC", "MC conversion on the current dmtgt.war with Master Data with Content",
     "2026-10-10 category changes (DB mode, all-locales fill, last-row label rule) re-tested for MNAO only.",
     "Schedule MC ja-JP with Master Data with Content (e-WD / WM sample).",
     "Categories written in DB mode (no Kapture login for categories); new categories present in all k_default_locales; "
     "no label UPDATED for keys Master Data Syncing already set; documents created/updated as before.",
     "MC Dev", "-", "High", "Pending"),
    ("T05", "Conversion re-test", "MME", "MME conversion on the current dmtgt.war with Master Data with Content",
     "Same as T04 - MME not re-tested since 2026-10-10.",
     "Schedule MME (e.g. en-UK / de-DE) with Master Data with Content.",
     "As T04 for the MME locale.",
     "MC Dev", "-", "High", "Pending"),
    ("T06", "Schedule - Only content check", "MC / MME", "Only content warning per selected locale",
     "Per-locale check verified for MNAO only (5694).",
     "1. Leave N rows in a locale that is NOT selected.\n2. Schedule MC / MME with Only content for another locale.",
     "No warning for rows of the other locale; warning when the selected locale has N rows, message shows the locale.",
     "MC Dev", "-", "Medium", "Pending"),
    # ---- B. Conversion items never exercised for lack of sample data
    ("T07", "Conversion - ESI", "MC", "ESI category mapping end to end",
     "Test esicat.txt used on MC Dev had 5-field lines (no ESI codes) - ESI path not proven for MC.",
     "Run an MC job whose esicat.txt carries ESI codes (8 fields, e.g. 04||041100||00).",
     "ESI TAXONOMY / MAPPED filled in the reports; documents tagged with the ESI categories in Kapture.",
     "MC Dev", "Business sample with ESI codes", "Medium", "Pending"),
    ("T08", "Preview", "MC", "WD HTML_MC / HTML5 file documents in Preview",
     "No sample content of that type so far.",
     "Convert a WD HTML_MC / HTML5 item, open the Preview page.",
     "Pages render with their own model folder; voltage / window.open links work.",
     "MC Dev", "Business sample", "Low", "Pending"),
    ("T09", "Preview", "All", "PDF inline display in Chrome",
     "Proven offline only.",
     "Open a PDF document in the Preview page in Chrome.",
     "PDF shows inline in the right pane.",
     "MC Dev", "-", "Low", "Pending"),
    ("T10", "Conversion - WD", "MME", "WD voltage links resolve",
     "84 links unresolved in 5617 because their targets were not loaded.",
     "Load NewM_CX-5_km/WM/FL0000/E-WZE-69P1_XML, re-run the WD item.",
     "Voltage links resolve; no 'unresolved' lines for them.",
     "MC Dev", "Target content loaded", "Low", "Pending"),
    ("T11", "Conversion - WM", "MME", "BT-50 WM full volume (5,214 docs)",
     "Only partial volumes run so far.",
     "Schedule the full BT-50 WM item.",
     "Job completes; timings and failures in line with earlier runs.",
     "MC Dev", "-", "Low", "Pending"),
    ("T12", "Conversion - WD", "MME", "window.js handling",
     "No MME sample with window.js yet.",
     "Run an MME item that carries window.js.",
     "window.js copied to OKAssets with /content/library prefix (MC rule).",
     "MC Dev", "Business sample", "Low", "Pending"),
    ("T13", "History", "MME / MNAO", "Late-day job visible in History",
     "Date filter fix (9 h timestamp shift) verified for MC only.",
     "Run an MME / MNAO job late in the day, filter History by today.",
     "Job listed.",
     "MC Dev", "-", "Low", "Pending"),
    # ---- C. Environment-dependent
    ("T14", "File behaviour", "All", "Re-check file-based behaviour after Syncthing",
     "rsync on /s3kapture brings deleted files back (preview.json, preview folders, reports, OKAssets).",
     "After Infra moves to Syncthing: run a delete job and a publish.",
     "Deleted preview docs stay gone from the earlier job's Preview page; report xlsx stay deleted after zipping; "
     "OKAssets deletes stay deleted. DB results unchanged (Delete counts, Kapture unpublished, preview_docs DELETED).",
     "MC Dev", "Infra: Syncthing live", "Medium", "Blocked"),
    # ---- D. Automation (Facelift) - built 2026-10-11 (option B load type, no schema change), not yet run
    ("T15", "Automation (Facelift)", "MC", "IPM Facelift VIN Update end to end - Only Facelift",
     "Built 2026-10-11: InfoManager replaced by Kapture (read latest article, swap only the VIN categories, update, publish).",
     "1. Put a ja_JP facelift Excel (11 columns) in Facelift_Content.\n2. Validate Data, Identify, then select items and schedule "
     "with Only Facelift.",
     "Impacted SM / OSM / WD docs found; in Kapture only the old VIN categories are replaced by the new ones (other categories, "
     "user groups, views, owner unchanged); gms3_vc_japan_vin_details rows carry the new VIN + carline / model names; "
     "processed / failed counts and reports correct.",
     "MC Dev", "-", "High", "Pending"),
    ("T16", "Automation (Facelift)", "MME", "IPM Facelift VIN Update end to end - Only Facelift",
     "As T15; MME view content table now lowercase (gms3_vc_mme_vin_dtl_<locale>).",
     "As T15 with an MME locale (e.g. en_UK) in the Excel.",
     "As T15 on gms3_vc_mme_vin_dtl_<locale>; gms3_mdm_vin_ml_mapping updated; MME VIN manual type mapping report in the zip.",
     "MC Dev", "-", "High", "Pending"),
    ("T17", "Automation (Facelift)", "MNAO", "IPM Facelift VIN Update end to end - Only Facelift",
     "As T15; legacy MNAO match kept (WMI / VDS / VIS, locale filter).",
     "As T15 with an en-US row in the Excel.",
     "As T15 on gms3_vc_vin_details (WMI / VDS / VIS of the matched rows replaced).",
     "MC Dev", "-", "High", "Pending"),
    ("T18", "Automation (Facelift)", "All", "Abort from History + conversion regression",
     "History abort of an automation job no longer clears gms3_dmt_el_gr_nm_dispord; Schedule.java dead code removed and its "
     "per-request fields moved to ThreadLocals (touches the conversion threads too).",
     "1. Start a conversion, then abort an automation job from History.\n2. Run one normal conversion per market.",
     "Automation job aborted; display-order staging of the running conversion untouched; conversions finish as before "
     "with the right user id on the schedule.",
     "MC Dev", "-", "Medium", "Pending"),
    # ---- E. QA
    ("T19", "QA regression", "MC", "Full DMT regression on QA",
     "DB changes applied on MC QA 2026-10-10 (Job 2 option 2); DMT app not yet tested there.",
     "Conversion (Master Data with Content + Only content), Publish, Delete, Master Data Syncing, Meta Data Validation, Automation.",
     "Same results as on MC Dev.",
     "MC QA", "dmtgt.war deployed on QA with QA properties", "High", "Pending"),
    ("T20", "QA regression", "MME", "Full DMT regression on QA", "As T19.", "As T19.", "As T19.",
     "MC QA", "As T19", "High", "Pending"),
    ("T21", "QA regression", "MNAO", "Full DMT regression on QA (incl. en-CA companion)", "As T19.", "As T19 + T01-T03.", "As T19.",
     "MC QA", "As T19", "High", "Pending"),
    # ---- D. Automation (Facelift) - continued (added 2026-10-11)
    ("T22", "Automation (Facelift)", "All", "Facelift folder read - no upload",
     "New 2026-10-11: every Excel of Facelift_Content (beside MDM) is read instead of an upload.",
     "1. No Facelift_Content folder / empty folder: open the IPM Facelift VIN Update page.\n"
     "2. Two Excel files with mixed locales + an open Excel (~$ lock file) in the folder: Validate Data.",
     "1. 'No Excel file found' message, no Validate button.\n2. Both files listed (lock file not), rows of both read and validated; "
     "messages name the file and row.",
     "MC Dev", "-", "High", "Pending"),
    ("T23", "Automation (Facelift)", "All", "Validation errors",
     "Validation now against MDM + Kapture (k_categories / k_default_locales) instead of OK_IM.",
     "Excel with: a blank cell, an unknown locale, an OLD VIN missing in Kapture, an OLD VIN missing in MDM, a NEW VIN in neither "
     "MDM / Kapture nor the ALL_VIN file. Then one with more than 10 errors.",
     "One error per case naming Kapture / MDM, row and file; Identify button not shown. More than 10 errors: text file link.",
     "MC Dev", "-", "High", "Pending"),
    ("T24", "Automation (Facelift)", "All", "New VIN only in the ALL_VIN master data file",
     "New 2026-10-11: a new VIN range missing in Kapture passes validation when the locale's ALL_VIN file carries it (D rows excluded).",
     "1. Excel row whose NEW VIN is only in the ALL_VIN file: Validate, Identify.\n2. Schedule with Only Facelift.\n"
     "3. Same range marked D in the ALL_VIN file: Validate.",
     "1. Validation passes.\n2. Refused with 'Only Facelift cannot be scheduled - n new VIN categories are not in Kapture yet'.\n"
     "3. Validation error for the new VIN.",
     "MC Dev", "-", "High", "Pending"),
    ("T25", "Automation (Facelift)", "All", "Master Data with Facelift",
     "New 2026-10-11: the master data of the item locales is loaded at the start of job B (like conversion); choice kept in dc_load_type.",
     "Continue T24: schedule the same items with Master Data with Facelift.",
     "Master data loaded first (new VIN category now in Kapture); documents updated with the new VIN and the refreshed model "
     "names; master data load / category reports inside the job zip; dc_load_type = MASTER_DATA_WITH_CONTENT.",
     "MC Dev", "T24", "High", "Pending"),
    ("T26", "Automation (Facelift)", "All", "Publish state kept",
     "D1: a document whose latest version is Published is updated and published again; a draft stays a draft.",
     "Include one document with a published latest version and one with a draft latest version in an item; run job B.",
     "k_article: the published one has a new live version carrying the new VIN; the draft one has a new draft, not published.",
     "MC Dev", "-", "High", "Pending"),
    ("T27", "Automation (Facelift)", "All", "Failed document keeps the old VIN",
     "Order per batch: Kapture first, view content only for documents Kapture accepted.",
     "Make one document fail in Kapture (e.g. locked / checked out by another user) and run job B.",
     "That document reported with KAPEXP003 / KAPEXP004 and counted as failed; its view content rows keep the OLD VIN; "
     "the other documents of the batch updated normally.",
     "MC Dev", "-", "Medium", "Pending"),
    ("T28", "Schedule page", "All", "Two users on the Schedule page at the same time",
     "Schedule.java per-request fields (messages, server path, user id) were shared servlet fields - now ThreadLocals.",
     "Two users (different iv-user) work on the Schedule page at the same time; each schedules a job.",
     "Each sees their own messages; each schedule carries its own user id.",
     "MC Dev", "-", "Medium", "Pending"),
    ("T29", "Regression", "All", "InfoManager removal - smoke of every screen",
     "2026-10-11: infomanager package, build-lib/infomanager and the IM properties removed (all dead code).",
     "Deploy the new dmtgt.war; open every screen; run one Master Data Syncing and one Meta Data Validation.",
     "No ClassNotFound / NoClassDefFound in the catalina log; screens and jobs work as before.",
     "MC Dev", "-", "High", "Pending"),
    ("T30", "Schedule - Conversion Go", "All", "Facelift run confirmation before the conversion records are listed",
     "New 2026-10-11: order with all three folders = MDM, then Facelift_Content (IPM Facelift VIN Update), then GMS3_Content.",
     "1. Excel in Facelift_Content: select conversion filters, click Go, answer No.\n2. Click Go again, answer Yes.\n"
     "3. No Excel in Facelift_Content (or no folder): click Go.",
     "1. Dialog 'Facelift Content found'; No = no records listed, nothing can be scheduled.\n2. Records listed as before.\n"
     "3. No dialog, records listed directly.",
     "MC Dev", "-", "High", "Pending"),
    ("T31", "Automation (Facelift)", "All", "Old VIN ranges made inactive in MDM after a successful job",
     "New 2026-10-12: a successful job sets mdm_vin_flag = deleted (sync status cleared) for its old VIN ranges; Kapture untouched.",
     "1. Run a job that ends Success.\n2. Run a job with at least one failed document.",
     "1. OLD_VIN_DEACTIVATION_REPORT.xlsx in the zip (DEACTIVATED, MDM locale shown); the old range no longer active in MDM; "
     "old VIN category still in Kapture.\n2. No deactivation, no report; log 'the old VIN ranges stay active in MDM'.",
     "MC Dev", "-", "High", "Pending"),
    ("T32", "Automation (Facelift)", "All", "Identification job after its items are scheduled - count, remark, View",
     "New 2026-10-12: the count is kept (legacy reduced it to 0) with a remark naming the update job; View only while an item is left.",
     "1. Identify, schedule ALL its items in an update job, refresh the job list during and after the job.\n"
     "2. Identify with 2+ items, schedule ONE, refresh.",
     "1. TOTAL DOCUMENTS unchanged; REMARKS '<n> documents (<locale> <channel> <model>) processed by MGSS_MDSYNC_...'; no View link. "
     "A View clicked on a stale page shows 'No item left to schedule...'.\n2. View still shown, lists only the remaining item.",
     "MC Dev", "-", "Medium", "Pending"),
    # ---- F. Added 2026-10-12 (delete.txt rule + Settings page guide)
    ("T33", "Conversion - display order file", "MC", "delete.txt is ignored (not read as a display order file)",
     "Fixed 2026-10-12: delete.txt is no longer used; the display order check compared the name with 'delete' only, so delete.txt "
     "was read as a display order file. Now: starts with d, ends with .txt, not delete.txt (Schedule screen + conversion).",
     "1. Put a delete.txt beside vin.txt / d01.txt in a facelift folder.\n2. Go, select the row, Schedule; let the job finish.",
     "Row accepted (no 'cannot be selected'); display order rows saved = d01.txt lines only; DISPLAY_ORDER / CD reports list "
     "d01.txt / CD00.txt only; delete.txt nowhere in reports; job Success.",
     "MC Dev", "-", "Medium", "Pending"),
    ("T34", "Conversion - display order file", "MME / MNAO", "delete.txt is ignored (not read as a display order file)",
     "Same fix in StartMMEConversionImpl / StartConversionMNAOImpl; run on MC only (T33).",
     "As T33 with an MME and an MNAO item.",
     "As T33.",
     "MC Dev", "-", "Low", "Pending"),
    ("T35", "Settings page", "All", "Content Provision Guide + Download PDF",
     "New 2026-10-12: 10-tab guide under the network path; PDF (cover + contents with page numbers, page footer) shipped in the war.",
     "1. Open Settings, click through all tabs, reload (last tab kept).\n2. Click Download PDF.",
     "Tabs switch, content readable; PDF downloads with contents, page numbers, links and bookmarks.",
     "Local + MC Dev", "-", "Low", "Pending"),
    ("T36", "Conversion - library paths sent to Kapture", "MC / MME / MNAO", "/content/ prefix on every /library/ path DMT sends to Kapture",
     "New 2026-10-13 (same rule as Job 10): kapture.library.path.context=/content. KaptureArticleBuilder prefixes "
     "HTML5_UPLOAD_DIRECTORY, DJVU_FILE_LOCATION, OASIS_FILE_LOCATION, OEM_FILE_LOCATION and every src / href starting "
     "/library/ in the HTML (img, a, link, embed). DMT tables + preview keep /library/.",
     "1. Deploy the war with kapture.library.path.context=/content in application.properties.\n"
     "2. Run one conversion per market: MC SM (with images + a PDF link), MME WD HTML5, MNAO WD HTML5 / DJVU (new layout), "
     "an OSM item, and an MC OEM item.\n"
     "3. In Kapture open each document (or /latest-article) and check the path attributes and the img src / a href.\n"
     "4. Open the documents in the Kapture viewer.",
     "Every library path reads /content/library/MAZDA/...; WD MNAO paths keep .../<model>/<manual type>/<facelift>/<material>/...; "
     "nothing is /content/content/; links to other documents (index?page=...) unchanged; images, PDFs and HTML5 / DJVU open.",
     "MC Dev", "-", "High", "Pending"),
    ("T37", "Conversion - display order / CD cleanup", "MC", "Documents no longer in the display order (or CD file) lose their view content (CD data)",
     "New 2026-10-13/14: for the vin.txt material folders of the item, ACTIVE documents the display order files do not list lose "
     "their view content (CD: CD data). A document converted again lost it already in the document step - recorded "
     "(ClearedDisplayData) so it is reported too. Report: new OPERATION column (INSERT / DELETE). Preview NOT changed.",
     "1. Run the familia item with the full d01.txt.\n2. Remove some PDF and HTML lines from d01.txt and re-run.\n"
     "3. Remove one more PDF and one more HTML line and re-run.",
     "Run 2/3: the removed documents = DISPLAY_ORDER.xlsx rows OPERATION DELETE, STATUS SUCCESS, remark 'VIEW CONTENT DELETED - "
     "DOCUMENT IS NOT IN THE DISPLAY ORDER FILE.'; log 'cleared by the document step and not re-created'; no view content rows left; "
     "they stay in the preview WITHOUT tree; documents dropped in an earlier run are not reported again.",
     "MC Dev", "-", "High", "Pending"),
    ("T38", "Conversion - view content status", "MC", "Listed but not converted document keeps its published view content status",
     "New 2026-10-14: VC_CONTENT_STATUS follows the DMT document's published flag (ViewContentStatus) instead of always Draft.",
     "1. Publish a familia PDF document.\n2. Remove its source file, keep its line in d01.txt, re-run the item.\n"
     "3. Query gms3_vc_japan_vin_details for the document.",
     "Document not in the preview, no DELETE row, INSERT SUCCESS in DISPLAY_ORDER.xlsx; its view content rows VC_CONTENT_STATUS = Published "
     "(Draft while it was unpublished).",
     "MC Dev", "-", "High", "Pending"),
    ("T39", "Conversion - display order / CD cleanup + view content status", "MME / MNAO",
     "T37 + T38 for MME (view content + CD) and MNAO (view content model year + VIN, no CD)",
     "Same code wired into StartMMEConversionImpl / StartConversionMNAOImpl and their DAOs; run on MC only (T37, T38).",
     "As T37 and T38 with an MME and an MNAO item.",
     "As T37 / T38 (MNAO: both gms3_vc_model_year_details and gms3_vc_vin_details).",
     "MC Dev", "T37, T38", "Medium", "Pending"),
]

EXTERNAL = [
    ("E01", "Kapture", "Bulk update rollbacks ('Plugin instructed the server to rollback') under load on MC Dev",
     "Re-run of the affected items shows 0 rollbacks once Kapture deploys the fix.", "Open"),
    ("E02", "Kapture", "4 locked WD docs (aborted run 5511) + 11 PDFs refused",
     "Re-run the MC item after Kapture clears them; docs created.", "Open"),
    ("E03", "Kapture", "Publish ~6.9 s/doc, bulk update ~2.3 s/doc; is parallel publish safe?",
     "Answer from Kapture; timing re-check if they change anything.", "Open"),
    ("E04", "Infra", "rsync -> Syncthing on /s3kapture", "Unblocks T14.", "Open"),
    ("E05", "Business data", "ENT <Revision Year/Month/Day> tags malformed (SAX error in log, CX-5 ENTs)",
     "Kept in the log on purpose as the marker; business fixes the tags.", "Open"),
    ("E06", "Business data", "Voltage JS paths without FL folder; Hybrid carline KB*C3 vs KBAC3; ESI 409 names",
     "Business corrects source content.", "Open"),
    ("E07", "Housekeeping", "Test category ZZDUPTEST01 on SC Dev (42 rows)",
     "Remove through the Kapture UI or Softclouds (no delete grant).", "Open"),
]

OPEN_DECISIONS = [
    ("O7", "Sanitizer: generic (Job 10 rules) vs legacy per-process cleaning", "Kept generic until decided."),
    ("D1-D8", "Automation decisions (publish state, Kapture read-back, DMT tables, vc_content_status, missing VIN category, channels, MNAO match, bulk)",
     "DECIDED 2026-10-11 and built: published stays published, read latest + swap VIN categories only, DMT tables untouched, "
     "missing new VIN = validation error unless in the ALL_VIN file, legacy MNAO match, set-based SQL. Load type = option B (dc_load_type)."),
    ("-", "Optional: skip unchanged documents on re-run (content hash); reuse the Kapture login", "User's call."),
]

# results of a test run: ID -> (Status, Job ID / Evidence, Tested on, Remarks) - override the Status of TESTS
RESULTS = {
    "T22": ("Passed", "Local deploy (UI + validation)", "2026-10-11", "Validated by the user on the local dmtgt.war."),
    "T23": ("Passed", "Local deploy (UI + validation)", "2026-10-11", "Validated by the user on the local dmtgt.war."),
    "T30": ("Passed", "Local deploy (UI)", "2026-10-11", "Validated by the user on the local dmtgt.war."),
    "T15": ("Passed", "MC Dev 5708 / 5715", "2026-10-12",
            "ja-JP WD CX-5 (KM), 820 docs, 0 failures both runs (5708 Master Data with Facelift, 5715 Only Facelift); VINs swapped in "
            "Kapture + gms3_vc_japan_vin_details (user checked WD25070 / WD28673); no VIN manual type report for MC (5715)."),
    "T26": ("Passed", "MC Dev 5715", "2026-10-12",
            "54 published docs re-published, report version 4.0 (read back from k_article), vc_content_status Published; drafts stay Draft "
            "(e.g. WD28673 0.3)."),
    "T32": ("In Progress", "MC Dev identification job 59 / 5715", "2026-10-12",
            "Step 1 OK (user): count stayed 820, remark shown, no View link. Step 2 (schedule one of two items) not run yet."),
    "T31": ("Passed", "MC Dev 5715", "2026-10-12",
            "Old range KM___KMY___100001199999 deactivated in MDM (ja-JP, 1 row); OLD_VIN_DEACTIVATION_REPORT.xlsx in the zip; Kapture untouched."),
    "T33": ("Passed", "MC Dev 5729", "2026-10-12",
            "ja-JP NewM_familia_van_ncp WM FL0001 with delete.txt in the folder: scheduled without error; 2,778 display order rows saved "
            "= d01.txt lines exactly (CD 2,778 from CD00.txt); DISPLAY_ORDER + CD reports 11/11 SUCCESS, d01.txt / CD00.txt only; "
            "TRANSACTION 11 SUCCESS, VIEW_CONTENT 22 SUCCESS; delete.txt in no report; 0 ERROR lines; preview 11 docs; job Success."),
    "T35": ("Passed", "Local + MC Dev deploy", "2026-10-12",
            "Validated by the user on both deployments: guide tabs and PDF download (v2 with contents page numbers and footer)."),
    "T37": ("Passed", "MC Dev 5736-5741", "2026-10-14",
            "familia NewM_familia_van_ncp WM FL0001: 5739 (option A war) reported SM231910, SM231920 (PDF) + SM231261 (HTML) as DELETE / "
            "SUCCESS ('cleared by the document step and not re-created'); 5740 / 5741 did not report them again (no view content left); "
            "preview after the revert shows dropped documents without tree. CD: no CD lines removed, 'all in the CD files'."),
    "T38": ("Passed", "MC Dev 5741 / 5742 / 5743", "2026-10-14",
            "SM231917 (J_DTSNAMELIST_4.pdf, source removed, kept in d01.txt): 5741 view content Draft (unpublished); published by 5742 "
            "from 5739 (0.6 -> 1.0); 5743 re-run: not in preview, INSERT SUCCESS, view content 2 rows Published (user query)."),
    "T39": ("Passed", "Accepted on the MC result (T37 / T38)", "2026-10-14",
            "Closed by the user without a separate run: MME and MNAO use the same shared code (UnlistedDisplayData, ClearedDisplayData, "
            "ViewContentStatus) - MC result accepted for both."),
}

CLOSED = [
    ("MC conversion + publish + delete", "MC", "5546 / 5553 / 5595 / 5602", "2026-10-05"),
    ("MME conversion + publish", "MME", "5603 / 5610 / 5617", "2026-10-05"),
    ("MNAO conversion + publish + delete", "MNAO", "5673 / 5666 / 5680", "2026-10-07"),
    ("Master Data Syncing (InfoManager-free, DB mode, all-locales fill)", "All", "899 / 906 / 913", "2026-10-10"),
    ("MNAO conversion re-test on the new war (Master Data with Content + Only content)", "MNAO", "5687 / 5694", "2026-10-10"),
    ("Label rule aligned between both screens", "MNAO", "5701 / 934", "2026-10-10"),
    ("D01 same refkey in two hierarchies - API = DB behaviour", "MNAO", "SC Dev API test", "2026-10-10"),
    ("Meta Data Validation (all types, all markets)", "All", "605-689", "2026-10-11"),
]

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HFILL = PatternFill("solid", fgColor="1F4E78")
HFONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HFILL, HFONT, Alignment(wrap_text=True, vertical="center"), BORDER
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
ws = wb.active
ws.title = "Pending Tests"
sheet(ws, HEAD, [t[:10] + RESULTS[t[0]] if t[0] in RESULTS else t + ("", "", "") for t in TESTS], [6, 22, 10, 34, 38, 42, 46, 10, 24, 9, 13, 18, 12, 30])
last = len(TESTS) + 1
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
ws.add_data_validation(dv)
dv.add("K2:K%d" % last)
for status, color in (("Passed", "C6EFCE"), ("Failed", "FFC7CE"), ("Blocked", "FFEB9C"), ("In Progress", "DDEBF7"), ("Not Applicable", "EDEDED")):
    ws.conditional_formatting.add("A2:N%d" % last, FormulaRule(formula=['$K2="%s"' % status], fill=PatternFill("solid", fgColor=color)))

ws2 = wb.create_sheet("Summary")
ws2.append(["Status", "Count"])
for s in STATUSES:
    ws2.append([s, '=COUNTIF(\'Pending Tests\'!K:K,"%s")' % s])
ws2.append(["Total", "=SUM(B2:B%d)" % (len(STATUSES) + 1)])
ws2.append([])
ws2.append(["Area", "Open (not Passed / Not Applicable)"])
start = ws2.max_row
areas = []
for t in TESTS:
    if t[1] not in areas:
        areas.append(t[1])
for a in areas:
    ws2.append([a, '=COUNTIFS(\'Pending Tests\'!B:B,"%s",\'Pending Tests\'!K:K,"<>Passed",\'Pending Tests\'!K:K,"<>Not Applicable")' % a])
for row in ws2.iter_rows():
    for c in row:
        c.border = BORDER
for r in (1, start):
    for c in ws2[r]:
        c.fill, c.font = HFILL, HFONT
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 34

ws3 = wb.create_sheet("External Dependencies")
sheet(ws3, ["ID", "Owner", "Item", "What closes it", "Status"], EXTERNAL, [6, 14, 60, 55, 10])
ws4 = wb.create_sheet("Open Decisions")
sheet(ws4, ["ID", "Decision", "Current state"], OPEN_DECISIONS, [8, 80, 45])
ws5 = wb.create_sheet("Closed (reference)")
sheet(ws5, ["Functionality", "Market", "Job IDs", "Closed on"], CLOSED, [70, 10, 26, 12])
ws6 = wb.create_sheet("Legend")
for line in [("Status", "Meaning"), ("Pending", "Not run yet"), ("In Progress", "Being run"), ("Passed", "Run, expected result seen - fill Job ID / Evidence + Tested on"),
             ("Failed", "Run, result differs - details in Remarks"), ("Blocked", "Cannot run until the 'Depends on' item is done"),
             ("Not Applicable", "Dropped by decision - reason in Remarks")]:
    ws6.append(list(line))
for c in ws6[1]:
    c.fill, c.font = HFILL, HFONT
ws6.column_dimensions["A"].width = 16
ws6.column_dimensions["B"].width = 70

wb.save(OUT)
print(OUT, len(TESTS), "tests")
